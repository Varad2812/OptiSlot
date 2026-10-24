"""The Revision Gauntlet: a deterministic Rev B built from ANY loaded project's Rev A.

    1. delete 1 tag                       4. add 1 new 12-pair cable
    2. change 1 standard DO to XYV        5. add 1 corrupt row (DO, instrument type UZV)
    3. add 1 tag on a spare pair          6. add 1 row missing its pair

Targets are chosen from the project's own data (never hard-coded tags), and the edits are made in the real
MDB workbook so the Rev B file can be downloaded and re-uploaded.
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass

import openpyxl

from engine.hardware import default_hw, listed_hw
from engine.pipeline import RunOutput


@dataclass
class Step:
    n: int
    action: str
    tag: str
    detail: str


def _mapped_cols(out: RunOutput) -> dict[str, int]:
    headers = out.table.headers
    return {m.field: headers.index(m.source) + 1 for m in out.mapping if m.source in headers}


def build(mdb: bytes, out: RunOutput) -> tuple[bytes, list[Step]]:
    wb = openpyxl.load_workbook(io.BytesIO(mdb))
    ws = wb[out.table.sheet] if out.table.sheet in wb.sheetnames else wb.worksheets[0]
    col = _mapped_cols(out)
    hdr_row = out.table.header_row
    steps: list[Step] = []
    slots = out.alloc.slots
    live = [s for s in out.signals if s.live and not s.quarantined and s.kind == "field"]
    by_cable: dict[str, list] = {}
    for s in live:
        by_cable.setdefault(s.cable, []).append(s)
    hw = {s.source_row: s for s in slots if s.status == "Field"}

    def setv(row: int, field: str, value: str) -> None:
        if field in col:
            ws.cell(row, col[field]).value = value

    def clone(src_row: int, **values: str) -> int:
        new = ws.max_row + 1
        for j in range(1, ws.max_column + 1):
            ws.cell(new, j).value = ws.cell(src_row, j).value
        for f, v in values.items():
            setv(new, f, v)
        return new

    used_tags = {s.tag_key for s in out.signals}

    def new_tag(template: str, n: int) -> str:
        m = re.match(r"^(.*?)(\d+)([^\d]*)$", template)
        base = f"{m.group(1)}9{n:04d}" if m else f"{template}-NEW{n}"
        while re.sub(r"\s+", "", base) in used_tags:
            n += 1
            base = f"{m.group(1)}9{n:04d}" if m else f"{template}-NEW{n}"
        used_tags.add(re.sub(r"\s+", "", base))
        return base

    # 1 — delete one analog input in the middle of a well-used cable
    victim = next((s for s in live if s.io_type == "AIR" and len(by_cable[s.cable]) >= 3 and s.pair and s.pair > 1), live[0])
    setv(victim.row, "status", "DELETED")
    steps.append(Step(1, "delete", victim.tag, f"Status → DELETED ({victim.cable} pair {victim.pair})"))

    # 2 — a standard DO becomes a solenoid (XYV): hardware swap in place
    hi_list = out.rules.params("hardware_map").get("lists", {}).get("high_current_do") or []
    std_hw = listed_hw(out.rules, "normal_do") or default_hw(out.rules)
    std_do = next((s for s in live if s.io_type == "DOR" and hi_list and hw.get(s.row) and std_hw
                   and (hw[s.row].charm, hw[s.row].tb) == std_hw), None)
    if std_do:
        solenoid = hi_list[0]       # the first high-current instrument type of the hardware map
        setv(std_do.row, "instrument_type", solenoid)
        setv(std_do.row, "instrument_desc", "Solenoid Valve")
        steps.append(Step(2, "modify", std_do.tag, f"Instrument type {std_do.instrument_type} → {solenoid} (solenoid)"))

    # 3 — a new tag on a spare pair of an existing cable (prefer a wired spare position)
    spare = next((s for s in slots if s.status == "Wired Spare" and s.cable in by_cable), None) or \
        next((s for s in slots if s.status == "Spare Slot" and s.cable in by_cable and s.pair), None)
    if spare:
        tmpl = next((m for m in by_cable[spare.cable] if m.io_type == spare.io_type), by_cable[spare.cable][0])
        tag = new_tag(tmpl.tag, 1)
        clone(tmpl.row, tag=tag, pair=f"Pair {spare.pair}", status="New")
        steps.append(Step(3, "add", tag, f"New tag on {spare.cable} pair {spare.pair} (was a {spare.status.lower()})"))

    # 4 — a brand-new 12-pair cable in the first controller domain (8 DI signals)
    first = live[0]
    unit = first.extra.get("unit", "") if first.extra.get("unit", "-") != "-" else "NEW"
    cable = f"{unit}-CE-9001"
    while cable in by_cable:
        cable = cable[:-1] + str(int(cable[-1]) + 1)
    di = next((s for s in live if s.io_type == "DIR" and s.domain == first.domain), first)
    spec = "12PX1MM2/FR/IND&OS/ARM/NIS"
    new_rows = []
    for p in range(1, 9):
        tag = new_tag(di.tag, 10 + p)
        new_rows.append(clone(di.row, tag=tag, cable=cable, cable_spec=spec, pair=f"Pair {p}", status="New", group="-"))
    steps.append(Step(4, "add-cable", cable, f"New 12-pair cable {cable} with 8 DI signals"))

    # 5 — a corrupt DO row: instrument type UZV is not in the hardware map
    do_t = next((s for s in live if s.io_type == "DOR"), di)
    tag = new_tag(do_t.tag, 30)
    clone(do_t.row, tag=tag, cable=cable, cable_spec=spec, pair="Pair 10", status="New", io_type="DO",
          instrument_type="UZV", instrument_desc="Shutdown Valve", group="-")
    steps.append(Step(5, "corrupt", tag, f"DO with unknown instrument type UZV on {cable} pair 10"))

    # 6 — a row missing its pair number
    tmpl = by_cable[victim.cable][0]
    tag = new_tag(tmpl.tag, 40)
    clone(tmpl.row, tag=tag, pair="", status="New")
    steps.append(Step(6, "corrupt", tag, f"New tag on {tmpl.cable} with no pair number"))

    buf = io.BytesIO()
    wb.save(buf)
    _ = hdr_row, new_rows
    return buf.getvalue(), steps
