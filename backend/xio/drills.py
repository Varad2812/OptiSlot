"""Headless drill runner: every surprise-dataset file through the real Service (import → allocate → verify → export).

    uv run python -m xio.drills            # prints a table, writes nothing
    uv run python -m xio.drills --md PATH  # also writes the markdown report

A drill PASSES when: no crash · every live tag is allocated or sits in a TQ/quarantine · the verifier reports 0 violations
· the allocation export re-opens in openpyxl. File-specific expectations (README_drills.md) are checked on top.
"""
from __future__ import annotations

import argparse
import io
import os
import tempfile
import traceback
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
DRILLS = ROOT / "data" / "drills"
SAMPLE = ROOT / "data" / "Sample_IO_MDB_VIT_v2_1.xlsx"


@dataclass
class DrillResult:
    name: str
    files: list[str]
    crashed: str = ""
    live: int = 0
    allocated: int = 0
    explained: int = 0          # live tags allocated OR covered by a TQ / quarantine
    quarantined: int = 0
    tqs: int = 0
    violations: int = 0
    verdict: str = ""
    export_ok: bool = False
    checks: list[tuple[str, bool, str]] = field(default_factory=list)   # (expectation, ok, detail)
    extra: dict = field(default_factory=dict)

    @property
    def passed(self) -> bool:
        return (not self.crashed and self.violations == 0 and self.export_ok and self.explained == self.live
                and all(ok for _, ok, _ in self.checks))


def _fresh_service():
    from app.service import Service
    from store.db import Store
    tmp = Path(tempfile.mkdtemp(prefix="optislot-drill-")) / "drill.db"
    return Service(Store(str(tmp)))


def _summarise(svc, pid: str, r: DrillResult) -> None:
    st = svc.get(pid)
    c = st.current
    out = c.out
    live = [s for s in out.signals if s.live]
    placed = {s.tag_key for s in out.alloc.slots if s.status in ("Field", "Cabinet Alarm")}
    tq_tags = {t.replace(" ", "") for q in c.tqs for t in q.tags}
    tq_rows = {x for q in c.tqs for x in q.rows}
    quarantined = {s.row for s in live if s.quarantined}
    r.live = len(live)
    r.allocated = sum(1 for s in live if s.tag_key in placed and not s.quarantined)
    r.explained = sum(1 for s in live if (s.tag_key in placed and not s.quarantined) or s.row in quarantined
                      or s.tag_key in tq_tags or s.row in tq_rows)
    r.quarantined = len(out.quarantine)
    r.tqs = len(c.tqs)
    r.violations = c.report.violations
    r.verdict = c.report.verdict
    wb = openpyxl.load_workbook(io.BytesIO(c.xlsx))
    ws = wb["IO Assignment"]
    r.export_ok = ws.max_row == len(c.out.rows) + 1 and ws.max_column == 76 and (ws.max_row - 1) % 384 == 0
    r.extra["cats"] = sorted({t.category for t in c.tqs})
    r.extra["optimiser"] = c.optimiser.method


def _expect(r: DrillResult, what: str, ok: bool, detail: str = "") -> None:
    r.checks.append((what, bool(ok), detail))


def _cats(svc, pid):
    return [t.category for t in svc.get(pid).current.tqs]


def run_drill(name: str, files: list[Path], check) -> DrillResult:
    r = DrillResult(name=name, files=[f.name for f in files])
    try:
        svc = _fresh_service()
        pid = "drill"
        svc.create_project(pid, name, files[0].read_bytes(), files[0].name)
        svc.import_files(pid, [(f.read_bytes(), f.name) for f in files])
        _summarise(svc, pid, r)
        check(svc, pid, r)
    except Exception:
        r.crashed = traceback.format_exc().strip().splitlines()[-1]
    return r


def _sample_ref():
    from engine.pipeline import run
    return run(SAMPLE)


# ───────────── per-drill expectations (README_drills.md) ─────────────

def chk_a(svc, pid, r):
    st = svc.get(pid)
    c = st.current
    maps = {m.field: m for m in c.out.mapping}
    for f in ("tag", "io_type", "cable", "pair", "group", "cable_spec", "instrument_type"):
        _expect(r, f"auto-maps `{f}`", bool(maps[f].source), f"{maps[f].source} ({maps[f].confidence})")
    proj = [t for t in c.tqs if t.category == "Project setting"]
    _expect(r, "ONE project-level TQ for the defaults", len(proj) == 1, f"{len(proj)}: " + "; ".join(t.message[:70] for t in proj))
    _expect(r, "full allocation", r.allocated == r.live, f"{r.allocated}/{r.live}")
    _expect(r, "no row-level default TQs", not [t for t in c.tqs if t.category in ("Missing room", "Missing area")])


def chk_b(svc, pid, r):
    c = svc.get(pid).current
    cats = _cats(svc, pid)
    q_reasons = " | ".join(q.reason for q in c.out.quarantine)
    quarantined_cats = {q.category for q in c.tqs if q.rows and any(x.row in q.rows for x in c.out.quarantine)}
    for what, cat in (("pair `1:0` → quarantined + TQ", "Unreadable pair"), ("`Pair 99` on a 24-pair cable → quarantined + TQ", "Pair exceeds cable size"),
                      ("missing pair → quarantined + TQ", "Missing pair"), ("IO type `XX` → quarantined + TQ", "Unknown IO type")):
        _expect(r, what, cat in quarantined_cats, cat)
    _expect(r, "no-cable row quarantined", "Missing cable" in quarantined_cats, "Missing cable")
    _expect(r, "duplicate tag → TQ", any("uplicate" in x for x in cats), ", ".join(sorted(set(cats))))
    _expect(r, "`UZV` allocated + TQ", "Unknown DO type" in cats, "")
    lower = [s for s in c.out.signals if s.live and s.io_type == "AIR" and not s.quarantined]
    _expect(r, "lowercase `ai` normalised to AIR and allocated", bool(lower), "")
    _expect(r, "header found below 3 title rows", c.out.table.header_row >= 4, f"header row {c.out.table.header_row}")


def chk_c(svc, pid, r):
    c = svc.get(pid).current
    cats = _cats(svc, pid)
    _expect(r, "mixed IS/NON-IS cable → TQ", any(q.rules_in_conflict == ["r-is"] for q in c.tqs), ", ".join(sorted(set(cats))))
    _expect(r, "A+B on one cable → TQ naming the rules",
            any("r-ab" in q.rules_in_conflict or any("ab" in x.lower() for x in q.rules_in_conflict) for q in c.tqs), "")
    _expect(r, "unknown segregation `Z` → TQ", any(x == "Unrecognised segregation code" and "'Z'" in q.message for x, q in zip(cats, c.tqs)), ", ".join(sorted(set(cats))))
    _expect(r, "IS never shares a baseplate with NON-IS (verifier)", c.report.violations == 0, f"{c.report.violations} violations")


def chk_d(svc, pid, r):
    t = svc.get(pid).current.out.table
    _expect(r, "picks sheet `IO List`", t.sheet == "IO List", t.sheet)
    ref = _sample_ref()
    _expect(r, "same result as the sample", r.allocated == ref.kpis.allocated, f"{r.allocated} vs {ref.kpis.allocated}")


def chk_e12(svc, pid, r):
    c = svc.get(pid).current
    ref = _sample_ref()
    dup = [q for q in c.tqs if "uplicate" in q.category]
    dup_tags = {t for q in dup for t in q.tags}
    _expect(r, "3 duplicate tags reported as a TQ", len(dup_tags) == 3, f"{len(dup_tags)} tags in {len(dup)} TQ(s)")
    _expect(r, "not double-allocated (total = sample)", r.allocated == ref.kpis.allocated, f"{r.allocated} vs {ref.kpis.allocated}")


def chk_e3(svc, pid, r):
    ref = _sample_ref()
    _expect(r, "same result as the xlsx sample", r.allocated == ref.kpis.allocated and r.tqs == len(ref.tqs),
            f"{r.allocated}/{r.tqs} TQ vs {ref.kpis.allocated}/{len(ref.tqs)} TQ")
    c = svc.get(pid).current
    _expect(r, "identical addresses to the xlsx sample",
            {s.tag_key: s.id for s in c.out.alloc.slots if s.status == "Field"}
            == {s.tag_key: s.id for s in ref.alloc.slots if s.status == "Field"}, "")


def run_drill_f() -> DrillResult:
    r = DrillResult(name="drill_F revB", files=["drill_F_revB_manual.xlsx"])
    try:
        svc = _fresh_service()
        svc.create_project("drill", "F", SAMPLE.read_bytes(), SAMPLE.name)
        from app.revisions import apply_revision
        f = DRILLS / "drill_F_revB_manual.xlsx"
        imp = apply_revision(svc, "drill", f.read_bytes(), f.name)
        _summarise(svc, "drill", r)
        _expect(r, "0 unchanged signals moved", imp.moved == [], f"{len(imp.moved)} moved")
        _expect(r, "3 slots freed", len(imp.deleted) == 3, f"{len(imp.deleted)} deleted")
        _expect(r, "stability ≈ 97–99%", 96.5 <= imp.stability_pct <= 99.5, f"{imp.stability_pct}%")
        mod = [m for m in imp.modified if "SE4302T02" in m.reason or "XYV" in m.reason]
        _expect(r, "DO → XYV swaps in place to SE4302T02 + SE4502", bool(mod) and all(m.from_slot == m.to_slot for m in mod), f"{len(mod)}")
        cur = svc.get("drill").current
        xyv = [s for s in cur.out.alloc.slots if s.change_type == "Modified" and s.charm == "SE4302T02"]
        _expect(r, "swapped hardware is T02 + SE4502", bool(xyv) and all(s.tb == "SE4502" for s in xyv), "")
        added = {m.tag: m for m in imp.added}
        _expect(r, "4 + 2 = 6 signals added", len(added) == 6, f"{len(added)}")
        spare = [m for m in imp.added if "existing cable" in m.reason]
        _expect(r, "2 new tags land on their own pair of an existing cable", len(spare) == 2, f"{len(spare)}")
        newcab = [s for s in cur.out.alloc.slots if s.cable == "320-CE-0999" and s.status == "Field"]
        _expect(r, "new cable 320-CE-0999 has 4 signals placed", len(newcab) == 4, f"{len(newcab)}")
    except Exception:
        r.crashed = traceback.format_exc().strip().splitlines()[-1]
    return r


def run_all() -> list[DrillResult]:
    D = DRILLS
    return [
        run_drill("A renamed/reordered/no room", [D / "drill_A_renamed_reordered_no_room.xlsx"], chk_a),
        run_drill("B title rows/junk", [D / "drill_B_title_rows_blank_rows_junk.xlsx"], chk_b),
        run_drill("C IS/AB/unknown seg", [D / "drill_C_IS_rows_AB_conflict_unknown_seg.xlsx"], chk_c),
        run_drill("D multisheet", [D / "drill_D_multisheet.xlsx"], chk_d),
        run_drill("E1+E2 multi-file", [D / "drill_E1_part1.xlsx", D / "drill_E2_part2_overlap.xlsx"], chk_e12),
        run_drill("E3 csv", [D / "drill_E3_full.csv"], chk_e3),
        run_drill_f(),
    ]


def markdown(results: list[DrillResult]) -> str:
    L = ["| Drill | Files | Live | Allocated | Quarantined | TQs | Violations | Export | Result |", "|---|---|---:|---:|---:|---:|---:|---|---|"]
    for r in results:
        res = "CRASH: " + r.crashed if r.crashed else ("PASS" if r.passed else "FAIL")
        L.append(f"| {r.name} | {', '.join(r.files)} | {r.live} | {r.allocated} | {r.quarantined} | {r.tqs} | {r.violations} "
                 f"({r.verdict}) | {'opens' if r.export_ok else 'BAD'} | **{res}** |")
    L.append("")
    for r in results:
        L.append(f"### {r.name}")
        L.append("")
        L.append(f"TQ categories: {', '.join(r.extra.get('cats', [])) or '—'} · optimiser: {r.extra.get('optimiser', '—')}")
        L.append("")
        for what, ok, detail in r.checks:
            L.append(f"- {'✅' if ok else '❌'} {what}" + (f" — {detail}" if detail else ""))
        L.append("")
    return "\n".join(L)


def main() -> None:
    os.environ.setdefault("OPTISLOT_DB", str(Path(tempfile.mkdtemp()) / "x.db"))
    ap = argparse.ArgumentParser()
    ap.add_argument("--md")
    a = ap.parse_args()
    res = run_all()
    text = markdown(res)
    print(text)
    if a.md:
        Path(a.md).write_text(text)


if __name__ == "__main__":
    main()
