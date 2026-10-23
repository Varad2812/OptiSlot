"""Excel outputs. The allocation is written with openpyxl into a COPY of the blank template (never
DataFrame.to_excel), so formatting is identical to the customer's file on every row. Files are byte-for-byte
reproducible: fixed document timestamps and normalised zip entry dates."""
from __future__ import annotations

import datetime as dt
import io
import re
import zipfile
from copy import copy
from pathlib import Path
from typing import Iterable

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.ruleslang import catalog

ROOT = Path(__file__).resolve().parents[2]
TEMPLATE = ROOT / "data" / "Blank_IO_Allocation_Template.xlsx"
FIXED_TIME = dt.datetime(2026, 1, 1, 0, 0, 0)
_INT = re.compile(r"^(0|[1-9]\d{0,9})$")


def normalise_zip(data: bytes) -> bytes:
    zin = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            zi = zipfile.ZipInfo(info.filename, date_time=(2026, 1, 1, 0, 0, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = info.external_attr
            body = zin.read(info.filename)
            if info.filename == "docProps/core.xml":  # openpyxl stamps 'modified' with the save time
                body = re.sub(rb"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)", rb"\g<1>2026-01-01T00:00:00Z\g<2>", body)
            zout.writestr(zi, body)
    return out.getvalue()


def _save(wb: openpyxl.Workbook) -> bytes:
    wb.properties.created = FIXED_TIME
    wb.properties.modified = FIXED_TIME
    wb.properties.creator = "OptiSlot"
    wb.properties.lastModifiedBy = "OptiSlot"
    buf = io.BytesIO()
    wb.save(buf)
    return normalise_zip(buf.getvalue())


def _cell_value(v: str):
    v = "" if v is None else str(v)
    return int(v) if _INT.match(v) else v


def write_allocation(rows: list[dict[str, str]], template: Path = TEMPLATE) -> bytes:
    wb = openpyxl.load_workbook(template)
    ws = wb[catalog()["template"]["sheet"]] if catalog()["template"]["sheet"] in wb.sheetnames else wb.worksheets[0]
    cols = [c.value for c in ws[1]]
    pre = ws.max_row  # rows 2..pre are pre-formatted in the template
    styles = [copy(ws.cell(2, j)._style) for j in range(1, len(cols) + 1)]
    height = ws.row_dimensions[2].height or 15
    for i, r in enumerate(rows, start=2):
        beyond = i > pre
        for j, col in enumerate(cols, start=1):
            cell = ws.cell(i, j, _cell_value(r.get(col, "-")))
            if beyond:
                cell._style = copy(styles[j - 1])
        if beyond:
            ws.row_dimensions[i].height = height
    last = max(len(rows) + 1, 2)
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{last}"
    ws.freeze_panes = "D2"
    return _save(wb)


# ───────────────────────────── styled report workbooks ─────────────────────────────

HEAD_FILL = PatternFill("solid", fgColor="1F3864")
HEAD_FONT = Font(name="Arial", size=9, bold=True, color="FFFFFF")
BODY_FONT = Font(name="Arial", size=9)
THIN = Side(style="thin", color="000000")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
OK_FILL = PatternFill("solid", fgColor="D9F2E3")
BAD_FILL = PatternFill("solid", fgColor="FBE0E0")
WARN_FILL = PatternFill("solid", fgColor="FFF2CC")


def add_table(wb: openpyxl.Workbook, title: str, header: list[str], rows: Iterable[list], widths: list[int] | None = None,
              status_col: str | None = None) -> int:
    ws = wb.create_sheet(title[:31])
    ws.append(header)
    for j in range(1, len(header) + 1):
        c = ws.cell(1, j)
        c.fill, c.font, c.border = HEAD_FILL, HEAD_FONT, BORDER
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.row_dimensions[1].height = 30
    n = 0
    si = header.index(status_col) + 1 if status_col and status_col in header else None
    for r in rows:
        ws.append([_cell_value(v) if isinstance(v, str) else v for v in r])
        n += 1
        for j in range(1, len(header) + 1):
            c = ws.cell(n + 1, j)
            c.font, c.border = BODY_FONT, BORDER
            c.alignment = Alignment(vertical="top", wrap_text=j == len(header))
        if si:
            v = str(ws.cell(n + 1, si).value or "").upper()
            fill = OK_FILL if v in ("PASS", "OK", "CLOSED", "APPLIED") else BAD_FILL if v in ("FAIL", "SHORT", "OPEN") else WARN_FILL if v in ("WARN", "EXCEPTED") else None
            if fill:
                ws.cell(n + 1, si).fill = fill
    for j, w in enumerate(widths or [], start=1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.auto_filter.ref = f"A1:{get_column_letter(len(header))}{max(n + 1, 2)}"
    ws.freeze_panes = "A2"
    return n


def new_book() -> openpyxl.Workbook:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    return wb


def save_book(wb: openpyxl.Workbook) -> bytes:
    return _save(wb)


def write_allocation_safe(rows: list[dict[str, str]], template: Path = TEMPLATE) -> tuple[bytes, BaseException | None]:
    """write_allocation, but a template/styling failure falls back to a plain workbook with the same 76 columns."""
    try:
        return write_allocation(rows, template), None
    except Exception as e:
        return _write_plain(rows), e


def _write_plain(rows: list[dict[str, str]]) -> bytes:
    cols = list(catalog()["template"]["columns"])
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = catalog()["template"]["sheet"]
    ws.append(cols)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3864")
    for r in rows:
        ws.append([_cell_value(r.get(c, "-")) for c in cols])
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{max(len(rows) + 1, 2)}"
    ws.freeze_panes = "D2"
    return _save(wb)


def write_table(headers: list[str], rows: list[list[str]], sheet: str = "MDB") -> bytes:
    """A plain, deterministic workbook (used to store several merged MDB files as one)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet
    ws.append(headers)
    for r in rows:
        ws.append(list(r))
    return _save(wb)
