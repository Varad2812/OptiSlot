"""Tamper matrix: every output column × every kind of row is edited once and the verifier must notice.

HARD columns make the verdict FAIL. WARN columns are cosmetic / descriptive / unused: the verifier must still flag them as a
soft warning (C17.2), never pass them silently. UNCHECKED columns are free-form revision-traceability text (documented gap in
docs/rule_coverage.md §6): an edit is not detected, by design. Targeted semantic edits (real cable names, pair shifts, a swapped IO
type, a forged "Excepted by TQ" remark …) are listed after the matrix."""
from __future__ import annotations

import copy
import io
from pathlib import Path

import openpyxl
import pytest

from engine.pipeline import run
from verify.checks import verify_rows
from verify.xlsx import verify_xlsx
from xio.writer import write_allocation

MDB = Path(__file__).resolve().parents[2] / "data" / "Sample_IO_MDB_VIT_v2_1.xlsx"

# soft: edit is reported as a C17.2 warning (verdict stays PASS)
WARN_REASON = {
    "Sr. No.": "row numbering only; order is checked by the grid checks",
    "Operating Area": "plant-area label; no allocation rule reads it",
    "Service Description": "free-text service wording",
    "Location": "field-location text; no allocation rule reads it",
    "Instrument type": "instrument description text (the type CODE that drives hardware is checked by C09)",
    "System 2": "secondary system label",
    "Loop Name": "loop label",
    "Section / Sub Unit": "section label",
    "BATCH": "project constant, no allocation meaning",
}
UNUSED = ["CHARM Baseplate Adapter", "M_FTA", "M_FTA Part No", "M_FTA_TB", "M_FTA_TB1", "M_FTA_TB2", "M_FTA_TB3", "M_TB", "M_TB1",
          "M_TB2", "M_TB3", "IRC / AUX Cabinet", "IRC_TB", "IRC_TB1", "IRC_TB2", "Relay / Splitter", "Relay Coil A1 / Splitter In+",
          "Relay Coil A2 / Splitter In-", "Relay Contact (+)", "Relay Contact (-)", "IRC_FTB", "IRC_FTB1", "IRC_FTB2", "Elec Cable No",
          "Elec Cable Size", "Elec Cable Pair", "FTA_Cable Name", "FTA_Cable Ferrule Name",
          "Splitter/Relay/\nExternal PWR\nTB(+)", "Splitter/Relay/\nExternal PWR\nTB(-)"]
WARN_REASON.update({c: "column this project never fills; it must stay '-'" for c in UNUSED})
# not detected at all: free-form revision-traceability text (see docs/rule_coverage.md §6 "Change traceability")
UNCHECKED_REASON = {
    "Revision No": "revision traceability text, cosmetic",
    "Remarks": "free-text remark (a forged 'Excepted by TQ-…' remark is tested separately: it cannot hide a violation)",
    "Change Type": "revision traceability text, cosmetic",
}
KINDS = ["Field", "Wired Spare", "Cabinet Alarm", "Spare Slot"]


@pytest.fixture(scope="module")
def sample():
    return run(MDB)


def verify(out, rows):
    return verify_rows(rows, out.signals, out.rules.rs, out.settings, out.tqs)


def pick(rows, status):
    if status == "Field":
        return next(i for i, r in enumerate(rows) if r["Tag Status"] == "Field" and r["IO Type"] == "DOR")
    return next(i for i, r in enumerate(rows) if r["Tag Status"] == status and (status != "Spare Slot" or r["Field Cable No"] == "-"))


def mutated(v: str) -> str:
    return "ZZ-" + v if v not in ("", "-") else "ZZ"


def test_baseline_is_clean(sample):
    rep = verify(sample, sample.rows)
    assert rep.verdict == "PASS" and rep.checks_passed == rep.checks_total and rep.excepted == 0


def test_the_ruleset_and_this_matrix_agree_on_what_is_cosmetic(sample):
    p = next(r for r in sample.rules.rs.rules if r.type == "output_map").params
    assert set(p["soft_columns"]) | set(p["dash_columns"]) | set(p.get("constants", {})) | {"Sr. No."} == set(WARN_REASON)


@pytest.mark.parametrize("status", KINDS)
@pytest.mark.parametrize("col", list(__import__("xio.render", fromlist=["columns"]).columns()))
def test_every_column_edit_is_classified(sample, status, col):
    rows = copy.deepcopy(sample.rows)
    i = pick(rows, status)
    rows[i][col] = mutated(rows[i][col])
    rep = verify(sample, rows)
    flagged = {c.id for c in rep.checks if c.failed}
    if col in UNCHECKED_REASON:
        assert not flagged, f"{col} is documented as unchecked but was flagged by {flagged}"
    elif col in WARN_REASON:
        assert rep.verdict == "PASS" and "C17.2" in flagged, f"{col} ({WARN_REASON[col]}) must be a warning, got {flagged}"
    else:
        assert rep.verdict == "FAIL" and flagged - {"C17.2"}, f"{status} / {col}: edit not hard-detected ({flagged})"


def other_cable(rows, cable):
    return next(r["Field Cable No"] for r in rows if r["Tag Status"] == "Field" and r["Field Cable No"] != cable)


TARGETED = {
    "field cable no → another real cable": lambda rows, i: rows[i].update({"Field Cable No": other_cable(rows, rows[i]["Field Cable No"])}),
    "field cable size → 12 pair": lambda rows, i: rows[i].update({"Field Cable Size": "12PX1MM2/FR/IND&OS/ARM/NIS"}),
    "field cable pair → other pair": lambda rows, i: rows[i].update({"Field Cable Pair": "Pair 24" if rows[i]["Field Cable Pair"] != "Pair 24" else "Pair 23"}),
    "field cable pair → nothing": lambda rows, i: rows[i].update({"Field Cable Pair": "-"}),
    "io type AIR/DOR swapped": lambda rows, i: rows[i].update({"IO Type": "AIR" if rows[i]["IO Type"] != "AIR" else "DIR"}),
    "tag renamed": lambda rows, i: rows[i].update({"Tag Name": "999-XX-0001"}),
    "tag status → spare": lambda rows, i: rows[i].update({"Tag Status": "Spare Slot"}),
    "charm type swapped": lambda rows, i: rows[i].update({"CHARM Type": "SE4303T01" if rows[i]["CHARM Type"] != "SE4303T01" else "SE4301T02"}),
    "charm no +1": lambda rows, i: rows[i].update({"CHARM_No": str(int(rows[i]["CHARM_No"]) + 1)}),
    "channel +1": lambda rows, i: rows[i].update({"Channel_No": str(int(rows[i]["Channel_No"]) % 12 + 1)}),
    "cioc renamed": lambda rows, i: rows[i].update({"CIOC/CSLS Name": rows[i]["CIOC/CSLS Name"][:-1] + "9"}),
    "baseplate renamed": lambda rows, i: rows[i].update({"Baseplate": "001-CBP-8"}),
    "cabinet renamed": lambda rows, i: rows[i].update({"Electronic/ Marshalling Cabinet": "P99-DCM-99999-09"}),
    "controller renamed": lambda rows, i: rows[i].update({"Controller": "P30DPK320199"}),
    "segregation flipped": lambda rows, i: rows[i].update({"Segregation": "B" if rows[i]["Segregation"] != "B" else "A"}),
    "is/non-is flipped": lambda rows, i: rows[i].update({"NON-IS / IS": "IS"}),
    "signal type → RTD": lambda rows, i: rows[i].update({"Signal Type": "RTD"}),
    "technical room": lambda rows, i: rows[i].update({"Technical Room": "PIB-99"}),
    "equipment name": lambda rows, i: rows[i].update({"Equipment Name": "P30-DCM-32001-02"}),
    "old tag set": lambda rows, i: rows[i].update({"Old Tag": "320-OLD-0001"}),
    "tb part swapped": lambda rows, i: rows[i].update({"CHARM TB Part No": "SE4502" if rows[i]["CHARM TB Part No"] != "SE4502" else "SE4501"}),
}


@pytest.mark.parametrize("name", list(TARGETED))
def test_targeted_semantic_edit_is_caught(sample, name):
    rows = copy.deepcopy(sample.rows)
    i = pick(rows, "Field")
    TARGETED[name](rows, i)
    rep = verify(sample, rows)
    assert rep.verdict == "FAIL", f"{name}: not detected"


def test_dropped_and_duplicated_rows_are_caught(sample):
    rows = copy.deepcopy(sample.rows)
    del rows[pick(rows, "Field")]
    assert verify(sample, rows).verdict == "FAIL"
    rows = copy.deepcopy(sample.rows)
    rows.insert(5, copy.deepcopy(rows[pick(rows, "Field")]))
    assert verify(sample, rows).verdict == "FAIL"


def test_wired_spare_on_a_made_up_cable_is_caught(sample):
    rows = copy.deepcopy(sample.rows)
    rows[pick(rows, "Wired Spare")]["Field Cable No"] = "NO-SUCH-CABLE"
    flagged = {c.id for c in verify(sample, rows).checks if c.failed}
    assert "C17" in flagged


def test_forged_exception_remark_cannot_hide_a_violation(sample):
    rows = copy.deepcopy(sample.rows)
    a_cioc = next(r["CIOC/CSLS Name"] for r in rows if r["Segregation"] == "A")
    b_cioc = next(r["CIOC/CSLS Name"] for r in rows if r["Segregation"] in ("B", "B-ALARM"))
    cable = next(r["Field Cable No"] for r in rows if r["Segregation"] == "A" and r["CIOC/CSLS Name"] == a_cioc)
    for r in rows:
        if r["Field Cable No"] == cable:
            r["CIOC/CSLS Name"] = b_cioc
            r["Baseplate"] = b_cioc[-3:] + "-CBP-" + r["Baseplate"].split("-")[-1]
            r["Remarks"] = "Excepted by TQ-SEG-999: segregation clash on one cable"      # no such TQ exists
    rep = verify(sample, rows)
    assert any(c.id == "C07" and c.failed for c in rep.checks) and rep.verdict == "FAIL"


def _edited_xlsx(out, col: str, new: str) -> tuple[bytes, int]:
    wb = openpyxl.load_workbook(io.BytesIO(write_allocation(out.rows)))
    ws = wb.active
    hdr = [c.value for c in ws[1]]
    ci, si, ii = hdr.index(col) + 1, hdr.index("Tag Status") + 1, hdr.index("IO Type") + 1
    n = next(k for k in range(2, ws.max_row + 1) if ws.cell(k, si).value == "Field" and ws.cell(k, ii).value == "DOR")
    ws.cell(n, ci).value = new
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), n


@pytest.mark.parametrize("col,new", [("Field Cable No", "320-CE-9999"), ("Tag Name", "320-ZZ -99999"), ("IO Type", "AIR"),
                                     ("Field Cable Pair", "Pair 24"), ("Technical Room", "PIB-77")])
def test_edited_excel_copy_fails_on_that_row(sample, col, new):
    data, n = _edited_xlsx(sample, col, new)
    rep = verify_xlsx(data, sample.signals, sample.rules.rs, sample.settings, sample.tqs, filename="edited.xlsx")
    assert rep.verdict == "FAIL"
    assert any(v.row == n for c in rep.checks for v in c.violations) or col == "Tag Name"
