"""Verifier tests: round trip, reference, planted breakages, tampered file, reproducibility, formatting."""
from __future__ import annotations

import copy
import io
from pathlib import Path

import openpyxl
import pytest
from hypothesis import HealthCheck, given, settings, strategies as st

from app.schemas import ProjectSettings
from engine.pipeline import run
from tests.helpers import row, table
from verify.certificate import certify
from verify.checks import verify_rows
from verify.xlsx import verify_xlsx
from xio.reports import capacity_bom, revision_impact, rule_check_report, tq_register
from xio.writer import TEMPLATE, write_allocation

ROOT = Path(__file__).resolve().parents[2]
MDB = ROOT / "data" / "Sample_IO_MDB_VIT_v2_1.xlsx"
REF = ROOT / "data" / "Sample_IO_Allocation_VIT_v2_1.xlsx"


@pytest.fixture(scope="module")
def sample():
    out = run(MDB)
    return out, write_allocation(out.rows)


def verify(out, rows):
    return verify_rows(rows, out.signals, out.rules.rs, out.settings, out.tqs)


def failed(rep) -> set[str]:
    return {c.id for c in rep.checks if c.failed}


def test_in_memory_pass(sample):
    out, _ = sample
    rep = verify(out, out.rows)
    assert rep.verdict == "PASS" and rep.violations == 0 and rep.checks_passed == rep.checks_total


def test_round_trip_exported_file(sample):
    out, data = sample
    rep = verify_xlsx(data, out.signals, out.rules.rs, out.settings, out.tqs)
    assert rep.verdict == "PASS", [(c.id, c.violations[:2]) for c in rep.checks if c.failed]
    assert rep.checks_passed == rep.checks_total


def test_reference_workbook_runs_and_catches_the_trap(sample):
    out, _ = sample
    rep = verify_xlsx(REF, out.signals, out.rules.rs, out.settings, out.tqs)
    assert rep.checks_total >= 18
    c01 = next(c for c in rep.checks if c.id == "C01")
    msgs = " ".join(v.message for v in c01.violations)
    assert "440-PI   -20077-L" in msgs and "440-PI   -20076-L" in msgs


def test_planted_breakages_are_caught(sample):
    out, _ = sample
    rows = copy.deepcopy(out.rows)
    # 1. swap an A cable onto a B CIOC
    a_cioc = next(r["CIOC/CSLS Name"] for r in rows if r["Segregation"] == "A")
    b_cioc = next(r["CIOC/CSLS Name"] for r in rows if r["Segregation"] in ("B", "B-ALARM"))
    a_cable = next(r["Field Cable No"] for r in rows if r["Segregation"] == "A" and r["CIOC/CSLS Name"] == a_cioc)
    for r in rows:
        if r["Field Cable No"] == a_cable:
            r["CIOC/CSLS Name"] = b_cioc
            r["Baseplate"] = b_cioc[-3:] + "-CBP-" + r["Baseplate"].split("-")[-1]
    rep = verify(out, rows)
    assert "C07" in failed(rep)

    # 2. duplicate a slot
    rows = copy.deepcopy(out.rows)
    f = [i for i, r in enumerate(rows) if r["Tag Status"] == "Field"]
    for k in ("CIOC/CSLS Name", "Baseplate", "Channel_No", "CHARM_No"):
        rows[f[1]][k] = rows[f[0]][k]
    assert "C02" in failed(verify(out, rows))

    # 3. drop a tag
    rows = copy.deepcopy(out.rows)
    dropped = rows.pop(f[5])
    rep = verify(out, rows)
    assert "C01" in failed(rep) and "C13" in failed(rep)
    assert any(dropped["Tag Name"] in v.message for c in rep.checks if c.id == "C01" for v in c.violations)


def test_tampered_file_fails_on_that_row(sample):
    out, data = sample
    wb = openpyxl.load_workbook(io.BytesIO(data))
    ws = wb.active
    hdr = [c.value for c in ws[1]]
    ci, ti, si = hdr.index("CHARM Type") + 1, hdr.index("Tag Name") + 1, hdr.index("Tag Status") + 1
    n = next(i for i in range(2, ws.max_row + 1) if ws.cell(i, si).value == "Field" and ws.cell(i, ci).value == "SE4303T01")
    ws.cell(n, ci).value = "SE4304T01"
    buf = io.BytesIO()
    wb.save(buf)
    rep = verify_xlsx(buf.getvalue(), out.signals, out.rules.rs, out.settings, out.tqs, filename="tampered.xlsx")
    assert rep.verdict == "FAIL"
    viol = [(c.id, v.row, v.tag) for c in rep.checks for v in c.violations]
    assert viol == [("C09", n, ws.cell(n, ti).value)], viol


def test_reproducible_hashes(sample):
    out, data = sample
    out2 = run(MDB)
    data2 = write_allocation(out2.rows)
    rep = verify(out, out.rows)
    c1 = certify(out.table.sha256, out.rules.rs, data, rep)
    c2 = certify(out2.table.sha256, out2.rules.rs, data2, verify(out2, out2.rows))
    assert c1.output_sha256 == c2.output_sha256 and c1.run_id == c2.run_id and c1.ruleset_hash == c2.ruleset_hash


def test_formatting_identical_to_template_on_every_row(sample):
    _, data = sample
    t = openpyxl.load_workbook(TEMPLATE).active
    ws = openpyxl.load_workbook(io.BytesIO(data)).active
    assert [c.value for c in ws[1]] == [c.value for c in t[1]]
    assert ws.freeze_panes == "D2" and ws.auto_filter.ref == f"A1:BX{ws.max_row}"
    assert ws.max_row == 1153
    ref = [(t.cell(2, j).font.name, t.cell(2, j).font.sz, t.cell(2, j).border.left.style, t.cell(2, j).border.bottom.style)
           for j in range(1, 77)]
    for i in range(2, ws.max_row + 1):
        got = [(ws.cell(i, j).font.name, ws.cell(i, j).font.sz, ws.cell(i, j).border.left.style, ws.cell(i, j).border.bottom.style)
               for j in range(1, 77)]
        assert got == ref, f"row {i} formatting differs"
        assert (ws.row_dimensions[i].height or 15) == 15
    assert t.cell(1, 1).fill.fgColor.rgb == ws.cell(1, 1).fill.fgColor.rgb


def test_reports_build(sample):
    out, data = sample
    rep = verify(out, out.rows)
    cert = certify(out.table.sha256, out.rules.rs, data, rep)
    for blob, n in (rule_check_report(rep, cert), tq_register(out.tqs, out.quarantine), revision_impact(None),
                    capacity_bom(out.alloc.cabinets, out.alloc.slots, out.kpis)):
        wb = openpyxl.load_workbook(io.BytesIO(blob))
        assert wb.sheetnames and n >= 0
        assert wb.worksheets[0].cell(1, 1).fill.fgColor.rgb.endswith("1F3864")


# ───────────────────── allocator and verifier never disagree ─────────────────────

mix = st.lists(st.tuples(st.sampled_from([1, 2, 6, 12, 24]), st.sampled_from(["A", "B", "-", "C", "A-CONTROL"]),
                         st.sampled_from(["AI", "AO", "DI", "DO"]), st.floats(0.1, 1.0),
                         st.sampled_from([("PIB-30", "RA 320-01"), ("PIB-24", "RA 500-02")])), min_size=1, max_size=50)


@settings(max_examples=120, deadline=None, suppress_health_check=[HealthCheck.too_slow])
@given(mix, st.booleans())
def test_engine_and_verifier_agree(m, ab_clash):
    rows = []
    for c, (size, grp, io, fill, (room, area)) in enumerate(m):
        n = max(1, int(size * fill))
        for p in range(1, n + 1):
            g = "B" if (ab_clash and c == 0 and p == n and n > 1) else grp
            rows.append(row(f"T{c}-{p}", io, f"CB-{c:03d}", p, size, group=g, room=room, area=area,
                            inst="XYV" if io == "DO" else "PT"))
    out = run(b"", table=table(rows))
    rep = verify_rows(out.rows, out.signals, out.rules.rs, out.settings, out.tqs)
    assert rep.verdict == "PASS", [(c.id, c.violations[:3]) for c in rep.checks if c.failed]


junk = st.text(alphabet=st.characters(blacklist_categories=("Cs",)), max_size=10)


@settings(max_examples=150, deadline=None, suppress_health_check=[HealthCheck.too_slow])
@given(st.lists(st.fixed_dictionaries({
    "tag": st.one_of(junk, st.sampled_from(["", "T-1", "T-2"])), "io": st.sampled_from(["AI", "AO", "DI", "DO", "XX", ""]),
    "cable": st.one_of(junk, st.sampled_from(["C-1", "C-2", "Internal Wiring", ""])),
    "pair": st.one_of(junk, st.integers(-2, 30).map(lambda n: f"Pair {n}")),
    "spec": st.sampled_from(["24PX1MM2", "12PX", "06PX", "2PX", "1TX", "", "-", "zz"]),
    "group": st.sampled_from(["A", "B", "-", "", "B-ALARM"]), "status": st.sampled_from(["New", "DELETED", ""]),
    "inst": st.sampled_from(["XYV", "XA", "UZV", "PT", ""]),
    "equip": st.sampled_from(["-", "P30-DCM-32001-02", "P30-DCS-32001-01", "P30-DCM-32001-07"]),
}), max_size=40))
def test_junk_engine_and_verifier_agree(rows):
    t = table([row(r["tag"], r["io"], r["cable"], r["pair"], group=r["group"], status=r["status"], equipment=r["equip"],
                   spec=r["spec"], inst=r["inst"]) for r in rows])
    out = run(b"", table=t, settings=ProjectSettings(max_cabinets_per_domain=8))
    rep = verify_rows(out.rows, out.signals, out.rules.rs, out.settings, out.tqs)
    assert rep.verdict == "PASS", [(c.id, c.violations[:3]) for c in rep.checks if c.failed]
