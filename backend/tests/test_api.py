"""API tests: the real engine behind the frozen contract."""
from __future__ import annotations

import io
import zipfile
from pathlib import Path

import openpyxl
import pytest
from fastapi.testclient import TestClient

from app.main import app
from app.schemas import AllocationResult, ImportSummary, RuleCheckReport, RuleSet

ROOT = Path(__file__).resolve().parents[2]
MDB = ROOT / "data" / "Sample_IO_MDB_VIT_v2_1.xlsx"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


@pytest.fixture(scope="module")
def c():
    with TestClient(app) as client:
        yield client


def test_demo_loaded(c):
    projects = c.get("/api/projects").json()
    assert any(p["id"] == "demo" for p in projects)
    a = AllocationResult.model_validate(c.get("/api/projects/demo/allocation").json())
    assert a.kpis.allocated == 189 and a.report.verdict == "PASS" and a.report.source == "file"
    assert len(a.slots) == 1152 and a.certificate.violations == 0


def test_allocate_is_reproducible(c):
    a1 = AllocationResult.model_validate(c.post("/api/projects/demo/allocate").json())
    a2 = AllocationResult.model_validate(c.post("/api/projects/demo/allocate").json())
    assert a1.certificate.output_sha256 == a2.certificate.output_sha256 and a1.run_id == a2.run_id
    runs = c.get("/api/projects/demo/runs").json()
    assert len(runs) >= 2 and runs[0]["output_sha256"] == a1.certificate.output_sha256


def test_import_and_mapping(c):
    r = c.post("/api/projects/demo/import", files={"file": ("Sample.xlsx", MDB.read_bytes(), XLSX)})
    s = ImportSummary.model_validate(r.json())
    assert s.rows == 190 and s.live_rows == 189 and len(s.preview) == 10
    tag = next(m for m in s.mapping if m.field == "tag")
    assert tag.source == "Tag Number" and tag.confidence == 1.0
    # override: map 'section' to P&ID, then back
    mapping = [m.model_dump() for m in s.mapping]
    for m in mapping:
        if m["field"] == "section":
            m["source"] = "P&ID"
    s2 = ImportSummary.model_validate(c.post("/api/projects/demo/mapping", json=mapping).json())
    assert next(m for m in s2.mapping if m.field == "section").source == "P&ID"
    for m in mapping:
        if m["field"] == "section":
            m["source"] = None
    c.post("/api/projects/demo/mapping", json=mapping)


def test_bad_upload_is_a_clean_error(c):
    r = c.post("/api/projects/demo/import", files={"file": ("junk.xlsx", b"not a workbook", XLSX)})
    assert r.status_code == 422


def test_rules_versioned_and_reallocate(c):
    rs = RuleSet.model_validate(c.get("/api/projects/demo/rules").json())
    v0 = len(c.get("/api/projects/demo/rules/versions").json())
    for r in rs.rules:
        if r.type == "wired_spare_pct":
            r.params["pct"] = 20
    a = AllocationResult.model_validate(c.put("/api/projects/demo/rules?why=test 20%", json=rs.model_dump()).json())
    assert a.kpis.wired_spares > 22 and a.report.verdict == "PASS"
    assert len(c.get("/api/projects/demo/rules/versions").json()) == v0 + 1
    assert c.get("/api/projects/demo/decisions").json()[0]["why"] == "test 20%"
    for r in rs.rules:
        if r.type == "wired_spare_pct":
            r.params["pct"] = 10
    a = AllocationResult.model_validate(c.put("/api/projects/demo/rules?why=back to 10%", json=rs.model_dump()).json())
    assert a.kpis.wired_spares == 22


def test_why_and_rows(c):
    a = AllocationResult.model_validate(c.get("/api/projects/demo/allocation").json())
    slot = next(s for s in a.slots if s.status == "Field")
    why = c.get(f"/api/projects/demo/slots/{slot.id}/why").json()
    assert len(why) >= 4
    rows = c.get("/api/projects/demo/rows").json()
    assert len(rows) == 1152 and len(rows[0]) == 76


def test_exports(c):
    files = c.get("/api/projects/demo/exports").json()
    assert {f["kind"] for f in files} == {"allocation", "rule_check", "tq_register", "revision_impact", "capacity_bom"}
    for f in files:
        r = c.get(f"/api/projects/demo/export/{f['kind']}")
        assert r.status_code == 200 and r.headers["content-type"] == XLSX
        openpyxl.load_workbook(io.BytesIO(r.content))
    z = zipfile.ZipFile(io.BytesIO(c.get("/api/projects/demo/export/all.zip").content))
    assert len(z.namelist()) == 5


def test_verify_file_tampered(c):
    data = c.get("/api/projects/demo/export/allocation").content
    ok = RuleCheckReport.model_validate(c.post("/api/verify-file", files={"file": ("ours.xlsx", data, XLSX)}).json())
    assert ok.verdict == "PASS"
    wb = openpyxl.load_workbook(io.BytesIO(data))
    ws = wb.active
    ws.cell(14, 25).value = "SE4304T01"
    buf = io.BytesIO()
    wb.save(buf)
    bad = RuleCheckReport.model_validate(c.post("/api/verify-file", files={"file": ("tampered.xlsx", buf.getvalue(), XLSX)}).json())
    assert bad.verdict == "FAIL" and any(v.row == 14 for ch in bad.checks for v in ch.violations)


def test_stress_endpoint(c):
    r = c.post("/api/projects/demo/stress?n=2000").json()
    assert r["signals"] >= 2000 and r["violations"] == 0


def test_unknown_project_404(c):
    assert c.get("/api/projects/nope/allocation").status_code == 404
