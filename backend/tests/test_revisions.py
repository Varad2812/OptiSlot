"""Revision tests: the Gauntlet, minimal moves, hard rules override freezing, persistence."""
from __future__ import annotations

import io
from pathlib import Path

import openpyxl
import pytest
from fastapi.testclient import TestClient

from app.schemas import AllocationResult, RevisionImpact
from engine.pipeline import run
from engine.revise import revise
from verify.checks import verify_rows
from xio.gauntlet import build

ROOT = Path(__file__).resolve().parents[2]
MDB = (ROOT / "data" / "Sample_IO_MDB_VIT_v2_1.xlsx").read_bytes()


@pytest.fixture(scope="module")
def gauntlet():
    a = run(MDB)
    revb, steps = build(MDB, a)
    b, imp = revise(revb, a)
    return a, b, imp, steps


def test_gauntlet_zero_violations(gauntlet):
    _, b, _, _ = gauntlet
    rep = verify_rows(b.rows, b.signals, b.rules.rs, b.settings, b.tqs)
    assert rep.verdict == "PASS" and rep.violations == 0


def test_gauntlet_expected_tqs(gauntlet):
    a, b, imp, _ = gauntlet
    cats = sorted(t.category for t in b.tqs if t.id in imp.new_tqs)
    assert cats == ["Missing pair", "Unknown DO type"]


def test_gauntlet_minimal_moves(gauntlet):
    a, b, imp, steps = gauntlet
    assert imp.moved == []
    deleted, modified = steps[0].tag, steps[1].tag
    before = {s.tag_key: s.id for s in a.alloc.slots if s.status in ("Field", "Cabinet Alarm")}
    after = {s.tag_key: s.id for s in b.alloc.slots if s.status in ("Field", "Cabinet Alarm")}
    for k, slot in before.items():
        if k in after:
            assert after[k] == slot, f"{k} moved"
    assert imp.unchanged_in_place == len(before) - 2  # one deleted, one modified in place
    assert imp.stability_pct == round(100 * (len(before) - 2) / len(before), 1)
    assert [m.tag for m in imp.deleted] == [deleted]


def test_gauntlet_hardware_swap_in_place(gauntlet):
    a, b, imp, steps = gauntlet
    key = steps[1].tag.replace(" ", "")
    sa = next(s for s in a.alloc.slots if s.tag_key == key)
    sb = next(s for s in b.alloc.slots if s.tag_key == key)
    assert sa.id == sb.id and (sa.charm, sa.tb) == ("SE4302T01", "SE4501") and (sb.charm, sb.tb) == ("SE4302T02", "SE4502")
    assert sb.change_type == "Modified" and sb.revision == "01"


def test_gauntlet_new_tag_on_former_spare(gauntlet):
    a, b, imp, steps = gauntlet
    key = steps[2].tag.replace(" ", "")
    sb = next(s for s in b.alloc.slots if s.tag_key == key)
    sa = next(s for s in a.alloc.slots if s.id == sb.id)
    assert sa.status == "Wired Spare" and sa.cable == sb.cable and sa.pair == sb.pair
    assert sb.change_type == "Added"


def test_gauntlet_deleted_slot_freed(gauntlet):
    a, b, imp, steps = gauntlet
    slot = next(s for s in b.alloc.slots if s.id == imp.deleted[0].from_slot)
    assert slot.status in ("Spare Slot", "Wired Spare") and slot.change_type == "Deleted"


def test_freezing_never_blocks_a_hard_rule():
    """Re-tag a whole A cable as B: it can no longer stay on its A CIOC, so it moves — with the rule as the reason."""
    a = run(MDB)
    wb = openpyxl.load_workbook(io.BytesIO(MDB))
    ws = wb.active
    hdr = [c.value for c in ws[1]]
    cab_col, grp_col = hdr.index("Multipair Cable No.") + 1, hdr.index("IO Grouping") + 1
    target = "320-CE-0035"
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, cab_col).value == target:
            ws.cell(r, grp_col).value = "B"
    buf = io.BytesIO()
    wb.save(buf)
    b, imp = revise(buf.getvalue(), a)
    moved_tags = {m.tag.replace(" ", "") for m in imp.moved}
    cable_tags = {s.tag_key for s in b.signals if s.cable == target and s.live}
    assert moved_tags == cable_tags and all("rule r-ab" in m.reason for m in imp.moved)
    rep = verify_rows(b.rows, b.signals, b.rules.rs, b.settings, b.tqs)
    assert rep.verdict == "PASS"
    others = {s.tag_key: s.id for s in a.alloc.slots if s.status == "Field" and s.cable != target}
    after = {s.tag_key: s.id for s in b.alloc.slots if s.status == "Field"}
    assert all(after[k] == v for k, v in others.items())


def test_api_gauntlet_switch_and_restart(tmp_path, monkeypatch):
    from app.service import Service
    from store.db import Store
    import app.main as main
    svc = Service(Store(tmp_path / "rev.db"))
    monkeypatch.setattr(main, "svc", svc)
    with TestClient(main.app) as c:
        imp = RevisionImpact.model_validate(c.post("/api/projects/demo/revisions/gauntlet").json())
        assert imp.stability_pct > 98 and imp.moved == []
        b = AllocationResult.model_validate(c.get("/api/projects/demo/allocation").json())
        assert b.revision == "B" and b.report.verdict == "PASS" and b.kpis.stability_pct == imp.stability_pct
        a = AllocationResult.model_validate(c.post("/api/projects/demo/revisions/switch?rev=A").json())
        assert a.revision == "A" and a.kpis.stability_pct is None
        b2 = AllocationResult.model_validate(c.post("/api/projects/demo/revisions/switch?rev=B").json())
        assert b2.certificate.output_sha256 == b.certificate.output_sha256
    # restart: a fresh service on the same database reloads Rev B with the same hash
    svc2 = Service(Store(tmp_path / "rev.db"))
    st = svc2.get("demo")
    assert st.project.revision == "B" and st.current.cert.output_sha256 == b.certificate.output_sha256
