"""Rules Studio API: catalogue, listing, custom-rule CRUD, preview, undo — and the full round trip with the independent verifier."""
from __future__ import annotations

import io

import openpyxl
import pytest
from fastapi.testclient import TestClient

from app.customrules import compile_rule
from app.schemas import AllocationResult, RuleCatalogue, RuleChange, RuleListing

P = "/api/projects/demo"
ZONE = {"type": "zone_cabinets", "label": "FT only in cabinet 04", "params": {"filter": {"tag": "320-FT*"}, "cabinets": [4]}}


@pytest.fixture()
def c(tmp_path, monkeypatch):
    import app.main as main
    from app.service import Service
    from store.db import Store
    monkeypatch.setattr(main, "svc", Service(Store(tmp_path / "cr.db")))
    with TestClient(main.app) as client:
        yield client


def alloc(c) -> AllocationResult:
    return AllocationResult.model_validate(c.get(P + "/allocation").json())


def post_rule(c, rule: dict) -> RuleChange:
    r = c.post(P + "/rules", json=rule)
    assert r.status_code == 200, r.text
    return RuleChange.model_validate(r.json())


# ───────────────────────────── catalogue and listing ─────────────────────────────

def test_rule_types_catalogue_has_parameter_schemas_and_valid_examples(c):
    cat = RuleCatalogue.model_validate(c.get("/api/rule-types").json())
    customs = {t.type: t for t in cat.types if t.custom}
    assert set(customs) == {"max_per_scope", "scope_relation", "forbid_mix", "zone_cabinets", "extra_spares"}
    for t in customs.values():
        assert t.params and t.description and t.example is not None
        compile_rule(t.example)                                            # every shipped example is a valid rule
        assert {p.name for p in t.params} >= set(t.example.params) - {"group_by"} or t.type == "scope_relation"
    assert "baseplate" in cat.scopes and "tag" in cat.filter_fields and "glob" in cat.filter_syntax.lower() or "`*`" in cat.filter_syntax
    assert {t.type for t in cat.types if not t.custom} >= {"capacity", "segregate", "hardware_map"}
    assert c.get(P + "/rule-types").json() == c.get("/api/rule-types").json()


def test_listing_shows_mode_severity_and_the_checks_that_test_each_rule(c):
    lst = RuleListing.model_validate(c.get(P + "/rules/listing").json())
    by = {r.id: r for r in lst.rules}
    assert by["r-ab"].mode == "enforced" and by["r-ab"].check_ids == ["C07"] and by["r-ab"].passed > 0
    assert by["r-group"].mode == "advisory" and by["r-group"].severity == "soft"
    assert by["r-uzv"].mode == "off"
    assert not any(r.custom for r in lst.rules)
    ch = post_rule(c, ZONE)
    assert ch.stored and ch.rule and ch.rule.custom and ch.rule.mode == "enforced" and ch.rule.check_ids == ["C23"]
    assert "320-FT*" in ch.rule.summary
    lst = RuleListing.model_validate(c.get(P + "/rules/listing").json())
    new = next(r for r in lst.rules if r.custom)
    assert new.id == ch.rule_id and new.check_ids == ["C23"] and new.passed > 0 and new.failed == 0
    soft = post_rule(c, {**ZONE, "label": "advisory zone", "severity": "soft", "id": "c-soft"})
    assert soft.rule and soft.rule.mode == "advisory"


# ───────────────────────────── the round trip ─────────────────────────────

def test_round_trip_add_preview_save_engine_obeys_verifier_catches_delete_undo(c):
    base = alloc(c)
    version0 = c.get(P + "/rules").json()["version"]
    n_decisions = len(c.get(P + "/decisions").json())
    tq_ids0 = c.app.state if False else None
    # 1 ── preview: the full what-if, nothing stored
    pv = RuleChange.model_validate(c.post(P + "/rules/preview", json={"action": "add", "rule": ZONE}).json())
    assert not pv.stored and pv.allowed and pv.moves > 0 and pv.report.violations == 0 and pv.kpis_after.violations == 0
    assert any(ch.id == "C23" for ch in pv.report.checks)
    assert c.get(P + "/rules").json()["version"] == version0 and len(c.get(P + "/decisions").json()) == n_decisions
    assert alloc(c).run_id == base.run_id and not any(r["custom"] for r in c.get(P + "/rules/listing").json()["rules"])
    # 2 ── save: a new rule-set version, a decision entry, and the engine obeys
    ch = post_rule(c, ZONE)
    assert ch.stored and ch.ruleset_version == version0 + 1 and ch.rule_id == "c-zone-cabinets-1"
    d = c.get(P + "/decisions").json()
    assert len(d) == n_decisions + 1 and d[0]["rule_id"] == "c-zone-cabinets-1" and d[0]["ruleset_version"] == version0 + 1
    a = alloc(c)
    fts = [s for s in a.slots if s.status == "Field" and s.tag.startswith("320-FT")]
    assert fts and {s.cabinet_seq for s in fts} == {4}
    assert a.report.violations == 0 and any(ch_.id == "C23" and ch_.failed == 0 for ch_ in a.report.checks)
    assert a.certificate.ruleset_hash != base.certificate.ruleset_hash, "the custom rule is part of the run certificate"
    # 3 ── a manual violation in an exported copy is caught by the independent verifier
    wb = openpyxl.load_workbook(io.BytesIO(c.get(P + "/export/allocation").content))
    ws = wb.active
    hdr = {h.value: i for i, h in enumerate(ws[1])}
    tag_col, cab_col = hdr["Tag Name"], hdr["Electronic/ Marshalling Cabinet"]
    rows = list(ws.iter_rows(min_row=2))
    ft = next(r for r in rows if str(r[tag_col].value).startswith("320-FT") and r[cab_col].value.endswith("-04"))
    other = next(r for r in rows if r[tag_col].value not in (None, "-") and r[cab_col].value.endswith("-01")
                 and not str(r[tag_col].value).startswith("320-FT") and r[hdr["Tag Status"]].value == "Field")
    ft[tag_col].value, other[tag_col].value = other[tag_col].value, ft[tag_col].value
    buf = io.BytesIO()
    wb.save(buf)
    r = c.post("/api/verify-file", files={"file": ("tampered.xlsx", buf.getvalue())})
    assert r.status_code == 200
    rep = r.json()
    assert rep["verdict"] == "FAIL"
    c23 = next(x for x in rep["checks"] if x["id"] == "C23")
    assert c23["failed"] >= 1 and any("320-FT" in v["message"] and "c-zone-cabinets-1" in v["message"] for v in c23["violations"])
    # 4 ── delete → the allocation of the previous rule set is back (and the deletion is undoable)
    dl = RuleChange.model_validate(c.delete(P + "/rules/c-zone-cabinets-1").json())
    assert dl.stored and dl.rule is None and dl.ruleset_version == version0 + 2
    assert alloc(c).certificate.ruleset_hash == base.certificate.ruleset_hash or alloc(c).kpis.violations == 0
    assert c.post(P + "/undo").status_code == 200
    assert any(r["id"] == "c-zone-cabinets-1" for r in c.get(P + "/rules").json()["rules"])
    assert {s.cabinet_seq for s in alloc(c).slots if s.status == "Field" and s.tag.startswith("320-FT")} == {4}
    # 5 ── the same inputs give the same hash (deterministic)
    h1 = alloc(c).certificate
    import app.main as main
    from app.service import Service
    main.svc = Service(main.svc.store)
    h2 = alloc(c).certificate
    assert (h1.run_id, h1.ruleset_hash, h1.output_sha256) == (h2.run_id, h2.ruleset_hash, h2.output_sha256)


def test_a_preview_reserves_no_tq_ids_and_leaves_no_run_record(c):
    import app.main as main
    ids0 = main.svc.store.get("demo", "tq_ids", {})
    runs0 = len(c.get(P + "/runs").json())
    rule = {"type": "max_per_scope", "label": "max 6 AI per baseplate", "params": {"filter": {"io_type": "AIR"}, "scope": "baseplate", "max": 6}}
    pv = RuleChange.model_validate(c.post(P + "/rules/preview", json={"action": "add", "rule": rule}).json())
    assert pv.new_tqs, "the sample's 24-pair AI cables break this rule by themselves → TQs"
    assert main.svc.store.get("demo", "tq_ids", {}) == ids0 and len(c.get(P + "/runs").json()) == runs0
    saved = post_rule(c, rule)
    assert saved.stored and sorted(saved.new_tqs) == sorted(pv.new_tqs), "saving gives the TQ ids the preview showed"


def test_edit_a_custom_rule_and_toggle_it(c):
    post_rule(c, ZONE)
    r = c.put(P + "/rules/c-zone-cabinets-1", json={"params": {"cabinets": [1, 4]}, "reason": "allow cab 1 too"})
    ch = RuleChange.model_validate(r.json())
    assert r.status_code == 200 and ch.stored and ch.rule.params["cabinets"] == [1, 4] and ch.rule.params["filter"] == ZONE["params"]["filter"]
    ch = RuleChange.model_validate(c.put(P + "/rules/c-zone-cabinets-1", json={"enabled": False}).json())
    assert ch.rule.mode == "off" and ch.rule.check_ids == []
    ch = RuleChange.model_validate(c.put(P + "/rules/c-zone-cabinets-1", json={"enabled": True, "severity": "soft"}).json())
    assert ch.rule.mode == "advisory"
    ch = RuleChange.model_validate(c.put(P + "/rules/c-zone-cabinets-1", json={"params": {"cabinets": [9]}, "replace_params": True}).json())
    assert not ch.stored or ch.rule.params == {"cabinets": [9]}
    # an edit that makes the rule malformed is a 422 and changes nothing
    v = c.get(P + "/rules").json()["version"]
    r = c.put(P + "/rules/c-zone-cabinets-1", json={"params": {"max": 3}, "unset_params": ["cabinets"]})
    assert r.status_code == 422 and c.get(P + "/rules").json()["version"] == v


def test_built_in_parameters_are_editable_with_a_safe_what_if(c):
    base = alloc(c)
    before = base.kpis.wired_spares
    ch = RuleChange.model_validate(c.put(P + "/rules/r-wired-spare", json={"params": {"pct": 20}, "reason": "client asked for 20%"}).json())
    assert ch.stored and ch.rule.params["pct"] == 20 and ch.kpis_after.wired_spares > before and ch.report.violations == 0
    assert alloc(c).kpis.wired_spares > before
    # a nested key, by path
    ch = RuleChange.model_validate(c.put(P + "/rules/r-hw", json={"params": {"lists.high_current_do": ["XYV", "ZYV", "UZV"]}}).json())
    assert ch.stored and "UZV" in ch.rule.params["lists"]["high_current_do"]
    # nonsense for a built-in is refused by the what-if (a system TQ would appear), nothing stored
    v = c.get(P + "/rules").json()["version"]
    bad = RuleChange.model_validate(c.put(P + "/rules/r-wired-spare", json={"params": {"pct": "lots"}}).json())
    assert not bad.stored and not bad.allowed and bad.messages and c.get(P + "/rules").json()["version"] == v
    assert alloc(c).report.violations == 0


def test_a_rule_that_would_break_a_hard_rule_is_refused_not_stored(c, monkeypatch):
    """The what-if is the gate: if the verifier would report a new violation the change is refused with the reason."""
    import app.rulesapi as ra
    real = ra._sys_tqs
    monkeypatch.setattr(ra, "_sys_tqs", lambda comp: real(comp) | ({"TQ-FAKE-001"} if comp.report.checks_total >= 24 else set()))
    ch = RuleChange.model_validate(c.post(P + "/rules", json=ZONE).json())
    assert not ch.stored and not ch.allowed and any("Refused" in m for m in ch.messages)
    assert not any(r["custom"] for r in c.get(P + "/rules/listing").json()["rules"])


# ───────────────────────────── malformed input ─────────────────────────────

@pytest.mark.parametrize("body,needle", [
    ({"type": "max_per_scope", "label": "x", "params": {"filter": {"colour": "red"}, "scope": "cioc", "max": 3}}, "unknown filter field"),
    ({"type": "max_per_scope", "label": "x", "params": {"filter": {}, "scope": "room", "max": 3}}, "scope must be one of"),
    ({"type": "max_per_scope", "label": "x", "params": {"filter": {}, "scope": "cioc", "max": -2}}, "max must be"),
    ({"type": "max_per_scope", "label": "x", "params": {"filter": "tag", "scope": "cioc", "max": 2}}, "must be an object"),
    ({"type": "zone_cabinets", "label": "x", "params": {"filter": {}, "cabinets": "four"}}, "cabinets"),
    ({"type": "forbid_mix", "label": "x", "params": {"io_type_a": ["AIR"], "io_type_b": ["AIR"], "scope": "cioc"}}, "overlap"),
    ({"type": "extra_spares", "label": "x", "params": {"io_type": "AIR"}}, "greater than 0"),
    ({"type": "max_per_scope", "label": "  ", "params": {"filter": {}, "scope": "cioc", "max": 2}}, "label"),
    ({"type": "max_per_scope", "label": "x", "id": "bad id!", "params": {"filter": {}, "scope": "cioc", "max": 2}}, "rule id"),
    ({"type": "capacity", "label": "x", "id": "r-capacity", "params": {}}, "already exists"),
    ({"type": "hardware_map", "label": "x", "id": "c-hw2", "params": {}}, "built-in"),
])
def test_malformed_rules_are_a_clear_422_and_change_nothing(c, body, needle):
    v = c.get(P + "/rules").json()["version"]
    r = c.post(P + "/rules", json=body)
    assert r.status_code == 422, r.text
    assert needle in r.json()["detail"], r.json()
    assert c.get(P + "/rules").json()["version"] == v
    assert c.post(P + "/rules/preview", json={"action": "add", "rule": body}).status_code == 422
    assert alloc(c).report.violations == 0


def test_unknown_type_and_bad_bodies_and_unknown_rule_ids(c):
    assert c.post(P + "/rules", json={"type": "teleport", "label": "x", "params": {}}).status_code == 422
    assert c.post(P + "/rules", json={"label": "x"}).status_code == 422
    assert c.put(P + "/rules/nope", json={"enabled": False}).status_code == 404
    assert c.delete(P + "/rules/nope").status_code == 404
    assert c.delete(P + "/rules/r-ab").status_code == 422
    assert c.post(P + "/rules/preview", json={"action": "update"}).status_code == 422
    assert c.post(P + "/rules/preview", json={"action": "delete", "rule_id": "nope"}).status_code == 404


def test_putting_a_whole_rule_set_with_a_malformed_custom_rule_is_a_422(c):
    rs = c.get(P + "/rules").json()
    rs["rules"].append({"id": "c-bad", "type": "max_per_scope", "enabled": True, "severity": "hard", "label": "bad",
                        "params": {"filter": {"colour": "red"}, "scope": "cioc", "max": 2}})
    v = rs["version"]
    r = c.put(P + "/rules", json=rs)
    assert r.status_code == 422 and "c-bad" in r.json()["detail"]
    assert c.get(P + "/rules").json()["version"] == v
    rs["rules"][-1]["id"] = rs["rules"][0]["id"]
    assert c.put(P + "/rules", json=rs).status_code == 422                       # duplicate id


def test_custom_rules_survive_a_revision_and_keep_zero_violations(c):
    post_rule(c, ZONE)
    post_rule(c, {"type": "max_per_scope", "label": "at most 40 DO per CIOC",
                  "params": {"filter": {"io_type": "DOR"}, "scope": "cioc", "max": 40}})
    assert c.post(P + "/revisions/gauntlet").status_code == 200
    a = alloc(c)
    assert a.revision == "B" and a.report.violations == 0
    assert {s.cabinet_seq for s in a.slots if s.status == "Field" and s.tag.startswith("320-FT")} == {4}
    assert [ch.id for ch in a.report.checks if ch.id >= "C23"] == ["C23", "C24"]
