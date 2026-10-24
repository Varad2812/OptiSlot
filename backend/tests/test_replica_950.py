"""Regression: the sample MDB repeated 5x with a -Xn suffix on tag, cable and loop names must verify with 0 violations.

The suffix turns the literal cable 'Internal Wiring' into five distinct pseudo-cables ('Internal Wiring-X1'…). Their
cabinet-alarm rows span two rooms, and C10 used to pull them all into the majority room's cabinet while the engine
(correctly) put each alarm in its own cabinet. A cabinet alarm carries no cable block, so it has no cable home."""
from __future__ import annotations

from pathlib import Path

import openpyxl

from engine.pipeline import run
from tests.helpers import row, table
from verify.checks import verify_rows

SAMPLE = Path(__file__).resolve().parents[2] / "data" / "Sample_IO_MDB_VIT_v2_1.xlsx"
SUFFIXED = ("Tag Number", "TagNumber", "Multipair Cable No.", "Loop Name")


def verified(out):
    return verify_rows(out.rows, out.signals, out.rules.rs, out.settings, out.tqs)


def make_replica(dst: Path, copies: int = 5) -> Path:
    wb = openpyxl.load_workbook(SAMPLE)
    ws = wb["MDB-1"]
    head = [c.value for c in ws[1]]
    cols = [head.index(n) for n in SUFFIXED]
    data = [[c.value for c in r] for r in ws.iter_rows(min_row=2)]
    data = [r for r in data if any(v is not None for v in r)]
    ws.delete_rows(2, ws.max_row)
    for k in range(1, copies + 1):
        for r in data:
            r = list(r)
            for i in cols:
                if isinstance(r[i], str) and r[i].strip():
                    r[i] = f"{r[i]}-X{k}"
            ws.append(r)
    wb.save(dst)
    return dst


def test_950_row_replica_verifies_with_zero_violations(tmp_path):
    out = run(make_replica(tmp_path / "replica950.xlsx"))
    assert out.table.rows and len(out.table.rows) == 950
    rep = verified(out)
    assert rep.verdict == "PASS" and rep.violations == 0, [(c.id, v.row, v.message) for c in rep.checks for v in c.violations][:5]
    assert rep.checks_total == 22 and rep.checks_passed == 22
    assert out.kpis.allocated == out.kpis.signals
    assert out.optimiser and out.optimiser.adopted        # CP-SAT is no longer rejected by the verifier


def test_alarm_pseudo_cable_spanning_two_rooms_stays_in_each_rooms_cabinet():
    rows = [row(f"A-{p}", "AI", "CB-1", p, 12) for p in range(1, 13)]
    for i, (room, area) in enumerate([("PIB-30", "RA 320-01")] * 2 + [("PIB-24", "RA 500-02")]):
        rows.append(row(f"X-ALM-{i}", "DI", "Internal Wiring-X1", None, 1, room=room, area=area, equipment="P30-DCM-32001-01",
                        spec="-"))
    out = run(None, table=table(rows))
    rep = verified(out)
    assert rep.verdict == "PASS", [(c.id, v.message) for c in rep.checks for v in c.violations]
    assert next(c for c in rep.checks if c.id == "C10").failed == 0
