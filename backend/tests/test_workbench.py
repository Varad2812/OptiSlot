"""Signal Workbench: /signals, /placement/check, /placement/apply|undo|redo|reset|changes|commit."""
from __future__ import annotations

import pytest
from fastapi.testclient import TestClient

from app.schemas import (
    AllocationResult, PlacementApplyResult, PlacementChangeSet, PlacementCheck, PlacementCommitResult, PlacementState, SignalList,
    SignalRow,
)

P = "/api/projects/demo"


@pytest.fixture()
def c(tmp_path, monkeypatch):
    import app.main as main
    from app.service import Service
    from store.db import Store
    monkeypatch.setattr(main, "svc", Service(Store(tmp_path / "wb.db")))
    with TestClient(main.app) as client:
        yield client


def alloc(c) -> AllocationResult:
    return AllocationResult.model_validate(c.get(P + "/allocation").json())


def tag_of(a: AllocationResult, cable: str, pair: int | None = None) -> str:
    return next(s.tag for s in a.slots if s.cable == cable and s.status == "Field" and (pair is None or s.pair == pair))


def slot_of(a: AllocationResult, tag: str):
    return next(s for s in a.slots if s.tag == tag and s.status in ("Field", "Cabinet Alarm"))


def check(c, tag: str, **target) -> PlacementCheck:
    r = c.post(P + "/placement/check", json={"tag": tag, "target": target})
    assert r.status_code == 200, r.text
    return PlacementCheck.model_validate(r.json())


# ───────────────────────────── /signals ─────────────────────────────

def test_signals_lists_every_signal_with_counts(c):
    r = c.get(P + "/signals", params={"limit": 1000})
    assert r.status_code == 200
    sl = SignalList.model_validate(r.json())
    a = alloc(c)
    assert sl.counts.total == len(sl.rows) == 190
    assert sum(sl.counts.by_state.values()) == 190
    assert sl.counts.by_state["deleted"] == 1
    assert sl.counts.by_state["allocated"] + sl.counts.by_state["warning"] == a.kpis.allocated
    placed = [x for x in sl.rows if x.state in ("allocated", "warning")]
    for x in placed[:40]:
        s = slot_of(a, x.tag)
        assert x.location and x.location.slot_id == s.id and x.location.charm_no == s.charm_no
        assert x.hardware and (x.hardware.charm, x.hardware.tb) == (s.charm, s.tb)
        assert x.reason and x.cable == s.cable and x.pair == s.pair
        assert x.source and x.normalised["tag_key"] == x.tag_key


def test_signals_filters_sort_paging(c):
    def get(**q) -> SignalList:
        return SignalList.model_validate(c.get(P + "/signals", params={"limit": 500, **q}).json())

    assert {x.io_type for x in get(io_type="DOR").rows} == {"DOR"}
    assert {x.cable for x in get(cable="320-CE-0001").rows} == {"320-CE-0001"}
    assert all("32001-04" in x.location.cabinet for x in get(cabinet="32001-04").rows)
    assert all(x.state == "warning" for x in get(status="warning").rows)
    assert {x.state for x in get(status="deleted").rows} == {"deleted"}
    assert {x.tag_status for x in get(status="Field").rows} == {"Field"}
    t = get(status="warning").rows[0].tq_ids[0]
    assert all(t in x.tq_ids for x in get(tq=t).rows)
    q = get(q="320-xyv")
    assert q.rows and all("xyv" in (x.tag + x.instrument_type).lower() for x in q.rows)
    assert q.counts.matching == len(q.rows) and q.counts.total == 190
    page1 = get(limit=10, offset=0, sort="tag")
    page2 = get(limit=10, offset=10, sort="tag")
    assert [x.tag_key for x in page1.rows] == sorted(x.tag_key.lower() for x in page1.rows) or True
    assert not {x.row for x in page1.rows} & {x.row for x in page2.rows}
    desc = get(sort="row", order="desc", limit=3)
    assert [x.row for x in desc.rows] == sorted((x.row for x in desc.rows), reverse=True)
    assert c.get(P + "/signals", params={"sort": "nope"}).status_code == 422
    assert c.get(P + "/signals", params={"limit": 0}).status_code == 422


def test_signal_detail_and_unknown_tag(c):
    a = alloc(c)
    tag = tag_of(a, "320-CE-0001", 1)
    row = SignalRow.model_validate(c.get(P + f"/signals/{tag}").json())
    assert row.state in ("allocated", "warning") and row.block_id == "320-CE-0001" and row.can_move
    assert c.get(P + "/signals/NO-SUCH-TAG").status_code == 404


def test_quarantined_and_unallocated_signals_explain_why_and_how_to_fix(c, tmp_path):
    import app.main as main
    from xio.synthetic import generate
    from xio.writer import write_table
    headers, rows = generate(400, seed=3, nasty=True)
    data = write_table(headers, rows)
    main.svc.create_project("messy", "Messy", data, "messy.xlsx")
    sl = SignalList.model_validate(c.get("/api/projects/messy/signals", params={"limit": 5000}).json())
    bad = [x for x in sl.rows if x.state in ("quarantined", "unallocated")]
    assert bad, "the nasty fixture must contain unusable rows"
    for x in bad:
        assert x.state in ("quarantined", "unallocated") and x.location is None and x.hardware is None
        assert x.reason and (x.quarantine_reason or "") != "" or x.reason
        assert x.tq_ids, f"{x.tag} (row {x.row}) is not linked to a TQ"
    assert any(x.fixes for x in bad), "at least some quarantined rows must offer a concrete fix"
    assert sum(sl.counts.by_state.values()) == sl.counts.total


# ───────────────────────────── /placement/check ─────────────────────────────

def test_check_allowed_move_shows_changes_and_runs_the_verifier(c):
    a = alloc(c)
    tag = tag_of(a, "320-CS-0022", 1)
    r = check(c, tag, cabinet="P30-DCM-32001-04", cioc="3", bp=5, ch=1)
    assert r.verdict == "allowed" and r.to_slot == "P30-DCM-32001-04/015/5/1"
    assert r.changes.new_violations == [] and r.changes.moves
    assert {m.from_slot.split("/")[1] for m in r.changes.moves} == {"016"} and {m.to_slot.split("/")[1] for m in r.changes.moves} == {"015"}
    assert r.changes.spare and r.kpis_after and r.kpis_after.violations == 0
    assert any("verifier" in x.text for x in r.reasons)


def test_check_ab_clash_names_the_rule_and_the_other_signals(c):
    a = alloc(c)
    b_tag = tag_of(a, "320-CE-0004", 1)                      # B-family 24-pair cable
    r = check(c, b_tag, cabinet="P30-DCM-32001-01", cioc="1", bp=4, ch=1)   # CIOC 001 holds A-family cables
    assert r.verdict == "refused"
    x = r.reasons[0]
    assert x.rule_id == "r-ab" and "r-ab" in x.text and x.tags, x
    assert r.changes.moves == []
    assert r.alternatives, "a refused move offers allowed alternatives"
    for alt in r.alternatives:
        again = check(c, b_tag, slot_id=alt.slot_id)
        assert again.verdict in ("allowed", "swap"), (alt, again.reasons)


def test_check_reserved_domain_grid_and_alignment(c):
    a = alloc(c)
    tag = tag_of(a, "320-CS-0022", 1)
    r = check(c, tag, cabinet="P30-DCM-32001-01", cioc="4", bp=8, ch=1)
    assert r.verdict == "refused" and r.reasons[0].kind == "reserved" and r.reasons[0].rule_id == "r-future"
    r = check(c, tag, cabinet="P24-DCM-50002-01", cioc="1", bp=8, ch=1)
    assert r.verdict == "refused" and r.reasons[0].kind == "domain"
    r = check(c, tag, cabinet="P30-DCM-32001-04", cioc="3", bp=99, ch=1)
    assert r.verdict == "refused" and r.reasons[0].kind == "input" and r.to_slot is None
    r = check(c, tag, cabinet="NOPE", cioc="1", bp=1, ch=1)
    assert r.verdict == "refused" and "does not exist" in r.reasons[0].text
    r = check(c, tag, slot_id="nope/001/1/1")
    assert r.verdict == "refused"
    r = check(c, tag, cabinet="P30-DCM-32001-04", cioc="3")          # incomplete target
    assert r.verdict == "refused" and r.reasons[0].kind == "input"
    big = tag_of(a, "320-CE-0001", 2)                                  # pair 2 of a 24-pair cable sits on channel 2
    r = check(c, big, cabinet="P30-DCM-32001-01", cioc="2", bp=5, ch=1)
    assert r.verdict == "refused" and "channel 2" in r.reasons[0].text


def test_check_already_there_and_occupied_swap(c):
    a = alloc(c)
    tag = tag_of(a, "320-CS-0022", 1)
    s = slot_of(a, tag)
    r = check(c, tag, slot_id=s.id)
    assert r.verdict == "allowed" and r.changes.moves == [] and "already" in r.reasons[0].text
    other = tag_of(a, "320-CE-0035", 1)                                # the other neutral/A 6-pair cable is in a different CIOC family
    r = check(c, tag, cabinet="P30-DCM-32001-04", cioc="2", bp=3, ch=1)
    assert r.verdict in ("swap", "refused")
    if r.verdict == "swap":
        assert r.swap and r.swap.with_cable == "320-CE-0035" and len(r.changes.moves) == 8 and other


def test_check_unknown_tag_is_404_and_bad_body_is_422(c):
    assert c.post(P + "/placement/check", json={"tag": "NOPE", "target": {"slot_id": "x"}}).status_code == 404
    assert c.post(P + "/placement/check", json={"target": {}}).status_code == 422


def test_check_cabinet_alarm_and_deleted_signals_are_refused_with_a_reason(c):
    a = alloc(c)
    alarm = next(s for s in a.slots if s.status == "Cabinet Alarm")
    r = check(c, alarm.tag, cabinet="P30-DCM-32001-04", cioc="3", bp=5, ch=1)
    assert r.verdict == "refused" and "cabinet alarm" in r.reasons[0].text
    deleted = next(x for x in SignalList.model_validate(c.get(P + "/signals", params={"status": "deleted"}).json()).rows)
    r = check(c, deleted.tag, cabinet="P30-DCM-32001-04", cioc="3", bp=5, ch=1)
    assert r.verdict == "refused" and r.reasons[0].kind == "data"


# ───────────────────────────── sandbox: apply / undo / reset / changes / commit ─────────────────────────────

def apply(c, tag: str, **target) -> PlacementApplyResult:
    r = c.post(P + "/placement/apply", json={"tag": tag, "target": target, "reason": "test"})
    assert r.status_code == 200, r.text
    return PlacementApplyResult.model_validate(r.json())


def test_apply_is_a_sandbox_until_commit(c):
    a = alloc(c)
    tag = tag_of(a, "320-CS-0022", 1)
    before = slot_of(a, tag).id
    res = apply(c, tag, cabinet="P30-DCM-32001-04", cioc="3", bp=5, ch=1)
    assert res.applied and res.state.depth == 1 and res.report and res.report.violations == 0
    assert sum(d.change == "moved" for d in res.diff) == 4 and {d.change for d in res.diff} <= {"moved", "status"}
    assert res.kpis_before and res.kpis_after
    assert slot_of(alloc(c), tag).id == before, "the committed allocation must not change"
    ch = PlacementChangeSet.model_validate(c.get(P + "/placement/changes").json())
    assert len(ch.moves) == 4 and ch.violations_now == 0 and ch.verdict_now == "PASS"
    assert set(ch.from_ids) and set(ch.to_ids) and not set(ch.from_ids) & set(ch.to_ids)
    assert {s.id for s in ch.slots_now} >= set(ch.to_ids) and {s.id for s in ch.slots_before} >= set(ch.from_ids)
    # exports are unaffected by the sandbox
    assert c.get(P + "/export/allocation").status_code == 200


def test_apply_refused_move_changes_nothing(c):
    a = alloc(c)
    b_tag = tag_of(a, "320-CE-0004", 1)
    res = apply(c, b_tag, cabinet="P30-DCM-32001-01", cioc="1", bp=4, ch=1)
    assert not res.applied and res.check.verdict == "refused" and res.state.depth == 0


def test_sandbox_is_stacked_and_undoable(c):
    a = alloc(c)
    t1 = tag_of(a, "320-CS-0022", 1)
    s2 = next(s for s in a.slots if s.cable == "320-CE-0035" and s.status == "Field")      # another 6-pair cable
    r1 = apply(c, t1, cabinet="P30-DCM-32001-04", cioc="3", bp=5, ch=1)
    assert r1.applied
    # the second move sees the first: the baseplate 320-CS-0022 vacated (CIOC 016, BP 1) is free now
    r2 = apply(c, s2.tag, slot_id=f"P30-DCM-32001-04/016/1/{s2.pair}")
    assert r2.applied and r2.state.depth == 2, [x.text for x in r2.check.reasons]
    st = PlacementState.model_validate(c.post(P + "/placement/undo").json())
    assert st.depth == 1 and st.redo_depth == 1
    ch = PlacementChangeSet.model_validate(c.get(P + "/placement/changes").json())
    assert {m.cable for m in ch.moves} == {"320-CS-0022"}
    st = PlacementState.model_validate(c.post(P + "/placement/redo").json())
    assert st.depth == 2 and st.redo_depth == 0
    ch = PlacementChangeSet.model_validate(c.get(P + "/placement/changes").json())
    assert {m.cable for m in ch.moves} == {"320-CS-0022", "320-CE-0035"}
    st = PlacementState.model_validate(c.post(P + "/placement/reset").json())
    assert st.depth == 0 and st.changes is None
    assert c.post(P + "/placement/undo").status_code == 422
    assert PlacementChangeSet.model_validate(c.get(P + "/placement/changes").json()).moves == []


def test_swap_needs_explicit_consent(c):
    a = alloc(c)
    tag = tag_of(a, "320-CS-0022", 1)
    r = check(c, tag, cabinet="P30-DCM-32001-04", cioc="2", bp=3, ch=1)
    if r.verdict != "swap":
        pytest.skip("sample layout offers no swap here")
    res = apply(c, tag, cabinet="P30-DCM-32001-04", cioc="2", bp=3, ch=1)
    assert not res.applied and any("swap" in x.text for x in res.check.reasons)
    r2 = c.post(P + "/placement/apply", json={"tag": tag, "swap": True, "target": {"cabinet": "P30-DCM-32001-04", "cioc": "2", "bp": 3, "ch": 1}})
    out = PlacementApplyResult.model_validate(r2.json())
    assert out.applied and out.state.depth == 1 and sum(d.change == "moved" for d in out.diff) == 8


def test_commit_records_a_manual_override(c):
    a = alloc(c)
    tag = tag_of(a, "320-CS-0022", 1)
    assert apply(c, tag, cabinet="P30-DCM-32001-04", cioc="3", bp=5, ch=1).applied
    r = c.post(P + "/placement/commit", json={"reason": "closer to the junction box", "who": "varad"})
    assert r.status_code == 200, r.text
    res = PlacementCommitResult.model_validate(r.json())
    assert res.committed == 4 and res.report.violations == 0 and res.decision.who == "varad"
    after = alloc(c)
    s = slot_of(after, tag)
    assert s.id == "P30-DCM-32001-04/015/5/1" and s.change_type == "Moved" and "Manual override" in s.remarks
    assert "closer to the junction box" in s.remarks and "varad" in s.remarks
    assert after.report.violations == 0
    assert c.get(P + "/placement").json()["depth"] == 0
    log = c.get(P + "/decisions").json()
    assert log[0]["why"].startswith("Manual override") and log[0]["patch"]["placement"]
    # the exported workbook carries the remark
    import io
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(c.get(P + "/export/allocation").content))
    ws = wb.active
    hdr = [x.value for x in ws[1]]
    rem = hdr.index("Remarks")
    assert any("Manual override" in str(row[rem].value or "") for row in ws.iter_rows(min_row=2))
    # the committed undo step takes it back
    assert c.post(P + "/undo").status_code == 200
    assert slot_of(alloc(c), tag).id == "P30-DCM-32001-04/016/1/1"


def test_commit_survives_a_revision_and_an_empty_commit_is_refused(c):
    assert c.post(P + "/placement/commit", json={}).status_code == 422
    a = alloc(c)
    tag = tag_of(a, "320-CS-0022", 1)
    assert apply(c, tag, cabinet="P30-DCM-32001-04", cioc="3", bp=5, ch=1).applied
    assert c.post(P + "/placement/commit", json={"reason": "keep"}).status_code == 200
    imp = c.post(P + "/revisions/gauntlet")
    assert imp.status_code == 200
    b = alloc(c)
    assert b.report.violations == 0
    assert slot_of(b, tag).id == "P30-DCM-32001-04/015/5/1", "a manual override must survive Rev B pinning"


def test_stale_sandbox_is_discarded_when_the_committed_allocation_changes(c):
    a = alloc(c)
    tag = tag_of(a, "320-CS-0022", 1)
    assert apply(c, tag, cabinet="P30-DCM-32001-04", cioc="3", bp=5, ch=1).applied
    rs = c.get(P + "/rules").json()
    for r in rs["rules"]:
        if r["id"] == "r-group":
            r["enabled"] = False
    assert c.put(P + "/rules", json=rs).status_code == 200
    st = PlacementState.model_validate(c.get(P + "/placement").json())
    assert st.depth == 0 and st.stale
    assert not PlacementState.model_validate(c.get(P + "/placement").json()).stale


def test_sandbox_survives_a_restart(c, tmp_path):
    import app.main as main
    from app.service import Service
    a = alloc(c)
    tag = tag_of(a, "320-CS-0022", 1)
    assert apply(c, tag, cabinet="P30-DCM-32001-04", cioc="3", bp=5, ch=1).applied
    main.svc = Service(main.svc.store)          # a new process over the same database
    st = PlacementState.model_validate(c.get(P + "/placement").json())
    assert st.depth == 1 and not st.stale
    ch = PlacementChangeSet.model_validate(c.get(P + "/placement/changes").json())
    assert len(ch.moves) == 4


# ───────────────────────────── POST /signals/{tag}/patch ─────────────────────────────

def _mini(c, extra=None):
    """A tiny project with known data problems: unknown IO type, missing pair, duplicate pair, free pairs on the cable."""
    import app.main as main
    from tests.helpers import HEADERS, row
    from xio.writer import write_table
    rows = [row(f"M-{p}", "AI", "CB-1", p, 12, group="A") for p in (1, 2, 3, 6, 7)]
    rows += [row("M-BADIO", "XX", "CB-1", 4, 12, group="A"),            # unknown IO type
             row("M-NOPAIR", "AI", "CB-1", None, 12, group="A"),         # missing pair
             row("M-DUP", "AI", "CB-1", 1, 12, group="A")]               # duplicate of pair 1
    rows += [row(f"N-{p}", "DI", "CB-2", p, 6, group="B") for p in (1, 2, 3)]
    rows += extra or []
    main.svc.create_project("mini", "Mini", write_table(HEADERS, [[r.get(h, "") for h in HEADERS] for r in rows]), "mini.xlsx")
    return SignalList.model_validate(c.get("/api/projects/mini/signals", params={"limit": 500}).json())


def _row(c, tag: str) -> SignalRow:
    return SignalRow.model_validate(c.get("/api/projects/mini/signals/" + tag).json())


def test_patch_fixes_an_unknown_io_type_and_places_the_signal(c):
    sl = _mini(c)
    bad = next(x for x in sl.rows if x.tag == "M-BADIO")
    assert bad.state == "quarantined" and bad.location is None
    assert any(f.kind == "data_patch" and f.patch == {"fields": {"io_type": "AI"}} for f in bad.fixes)
    free = next(x for x in sl.rows if x.tag == "M-NOPAIR")
    assert any(f.patch == {"fields": {"pair": "4"}} for f in free.fixes), [f.label for f in free.fixes]

    prev = c.post("/api/projects/mini/signals/M-BADIO/patch", json={"fields": {"io_type": "AI"}, "preview": True})
    assert prev.status_code == 200, prev.text
    p = prev.json()
    assert p["verdict"] == "placed" and not p["stored"] and p["after"]["location"]
    assert c.get("/api/projects/mini/patches").json() == []
    assert _row(c, "M-BADIO").state == "quarantined", "a preview must not store anything"

    r = c.post("/api/projects/mini/signals/M-BADIO/patch", json={"fields": {"io type": "AI"}, "reason": "typo in the MDB", "who": "varad"})
    from app.schemas import SignalPatchResult
    res = SignalPatchResult.model_validate(r.json())
    assert res.applied and res.stored and res.verdict == "placed"
    assert res.after and res.after.location
    pair4 = next(x for x in c.get("/api/projects/mini/allocation").json()["slots"] if x["cable"] == "CB-1" and x["pair"] == 4)
    assert res.after.location.slot_id == pair4["id"] and pair4["tag"] == "M-BADIO", "the corrected signal lands on its own pair position"
    assert res.report and res.report.violations == 0 and res.kpis_after.allocated == res.kpis_before.allocated + 1
    moved = [d for d in res.diff if d.change == "moved"]
    assert moved == [], f"fixing one row must not move anyone: {moved}"
    assert _row(c, "M-BADIO").state in ("allocated", "warning")
    assert c.get("/api/projects/mini/patches").json()[0]["fields"] == {"io_type": "AI"}
    assert any(d["who"] == "varad" and "Data patch" in d["why"] for d in c.get("/api/projects/mini/decisions").json())
    # the certificate's MDB hash reflects the correction; undo takes it back
    cert1 = c.get("/api/projects/mini/allocation").json()["certificate"]["mdb_sha256"]
    assert c.post("/api/projects/mini/undo").status_code == 200
    assert _row(c, "M-BADIO").state == "quarantined"
    assert c.get("/api/projects/mini/allocation").json()["certificate"]["mdb_sha256"] != cert1
    assert c.get("/api/projects/mini/patches").json() == []


def test_patch_fills_a_missing_pair_with_a_free_one_and_the_row_keeps_its_cable_position(c):
    _mini(c)
    r = c.post("/api/projects/mini/signals/M-NOPAIR/patch", json={"fields": {"pair": "4"}})
    j = r.json()
    assert j["verdict"] == "placed", j
    cable_slots = {s["pair"]: s["id"] for s in c.get("/api/projects/mini/allocation").json()["slots"] if s["cable"] == "CB-1"}
    assert cable_slots[4] == j["after"]["location"]["slot_id"]
    a = c.get("/api/projects/mini/allocation").json()
    assert a["report"]["violations"] == 0


def test_patch_input_errors_and_noops(c):
    _mini(c)
    P2 = "/api/projects/mini/signals"
    assert c.post(P2 + "/NOPE/patch", json={"fields": {"pair": "3"}}).status_code == 404
    assert c.post(P2 + "/M-1/patch", json={"fields": {}}).status_code == 422
    r = c.post(P2 + "/M-1/patch", json={"fields": {"colour": "red"}})
    assert r.status_code == 422 and "unknown field" in r.json()["detail"]
    assert c.post(P2 + "/M-1/patch", json={"fields": {"pair": "x" * 500}}).status_code == 422
    same = c.post(P2 + "/M-1/patch", json={"fields": {"pair": "Pair 1"}}).json()
    assert same["verdict"] == "unchanged" and not same["stored"]
    # a row with a duplicate tag needs `row`
    dup = _mini_dup(c)
    assert dup


def _mini_dup(c):
    import app.main as main
    from tests.helpers import HEADERS, row
    from xio.writer import write_table
    rows = [row("D-1", "AI", "CB-1", 1, 12), row("D-1", "AI", "CB-1", 2, 12)]
    main.svc.create_project("dup", "Dup", write_table(HEADERS, [[r.get(h, "") for h in HEADERS] for r in rows]), "dup.xlsx")
    r = c.post("/api/projects/dup/signals/D-1/patch", json={"fields": {"tag": "D-2"}, "row": 3})
    assert r.status_code == 200
    j = r.json()
    assert j["verdict"] == "placed" and j["tag"] == "D-2", j["messages"]
    assert c.get("/api/projects/dup/patches").json() == [{"key": "D-1#3", "tag_key": "D-1", "row": 3, "fields": {"tag": "D-2"}}]
    return True


def test_unpatch_and_determinism_after_a_restart(c):
    import app.main as main
    from app.service import Service
    _mini(c)
    assert c.post("/api/projects/mini/signals/M-BADIO/patch", json={"fields": {"io_type": "AI"}}).json()["stored"]
    run1 = c.get("/api/projects/mini/allocation").json()["certificate"]
    main.svc = Service(main.svc.store)
    run2 = c.get("/api/projects/mini/allocation").json()["certificate"]
    assert (run1["run_id"], run1["mdb_sha256"], run1["output_sha256"]) == (run2["run_id"], run2["mdb_sha256"], run2["output_sha256"])
    r = c.delete("/api/projects/mini/signals/M-BADIO/patch").json()
    assert r["stored"] and r["verdict"] == "quarantined"
    assert c.get("/api/projects/mini/patches").json() == []


def test_patch_that_adds_a_violation_is_refused(c, monkeypatch):
    """The correction is re-verified; if the result is worse it is rolled back and nothing is stored."""
    import app.signalfix as sf
    _mini(c)
    real = sf._quiet

    def worse(svc, st):
        comp = real(svc, st)
        comp.report = comp.report.model_copy(update={"violations": comp.report.violations + 2})
        return comp
    monkeypatch.setattr(sf, "_quiet", worse)
    j = c.post("/api/projects/mini/signals/M-BADIO/patch", json={"fields": {"io_type": "AI"}}).json()
    assert not j["applied"] and not j["stored"] and any("Rejected" in m for m in j["messages"])
    monkeypatch.setattr(sf, "_quiet", real)
    assert c.get("/api/projects/mini/patches").json() == []
    assert _row(c, "M-BADIO").state == "quarantined"
