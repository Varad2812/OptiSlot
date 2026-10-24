"""Revision chains A → B → C (each pinned against the previous), TQ Apply as a pinned re-allocation, TQ data patches."""
from __future__ import annotations

import io

import openpyxl
import pytest

from app.revisions import apply_revision, run_gauntlet, switch
from app.sandbox import apply_tq
from app.service import Service, rev_no
from store.db import Store
from xio.drills import DRILLS, SAMPLE

SAMPLE_BYTES = SAMPLE.read_bytes()


def fresh(tmp_path, mdb: bytes = SAMPLE_BYTES, name: str = SAMPLE.name) -> Service:
    svc = Service(Store(tmp_path / "t.db"))
    svc.create_project("p", "p", mdb, name)
    return svc


def slots_of(out) -> dict[str, str]:
    return {s.tag_key: s.id for s in out.alloc.slots if s.status in ("Field", "Cabinet Alarm")}


def edit(mdb: bytes, fn) -> bytes:
    wb = openpyxl.load_workbook(io.BytesIO(mdb))
    fn(wb.worksheets[0])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def header(ws) -> dict[str, int]:
    return {str(c.value): c.column for c in ws[1] if c.value}


# ───────────────────────────── revision numbers ─────────────────────────────

def test_revision_numbers_come_from_the_project_revision():
    assert [rev_no(x) for x in "ABCD"] == ["00", "01", "02", "03"]
    assert [rev_no(x, "R05") for x in "ABC"] == ["R05", "R06", "R07"]
    assert rev_no("C", "draft") == "02"


def test_revision_number_written_to_the_sheet_follows_settings(tmp_path):
    svc = fresh(tmp_path)
    st = svc.get("p")
    st.settings = st.settings.model_copy(update={"revision": "07"})
    svc.store.put("p", "settings", st.settings.model_dump())
    svc.recompute(st, record=False)
    imp = run_gauntlet(svc, "p")
    assert imp.to_rev == "B"
    revs = {s.revision for s in st.current.out.alloc.slots if s.change_type in ("Added", "Modified")}
    assert revs == {"08"}


# ───────────────────────────── chains ─────────────────────────────

@pytest.fixture(scope="module")
def chain(tmp_path_factory):
    svc = fresh(tmp_path_factory.mktemp("chain"))
    st = svc.get("p")
    a = slots_of(st.current.out)
    ib = run_gauntlet(svc, "p")
    b = slots_of(st.current.out)
    ic = run_gauntlet(svc, "p", rev="C")
    c = slots_of(st.current.out)
    return svc, st, a, b, c, ib, ic


def test_chain_has_three_revisions(chain):
    svc, st, *_ , ib, ic = chain
    assert st.project.revisions == ["A", "B", "C"] and st.project.revision == "C"
    assert (ib.from_rev, ib.to_rev, ic.from_rev, ic.to_rev) == ("A", "B", "B", "C")


def test_rev_c_is_pinned_against_rev_b_not_rev_a(chain):
    svc, st, a, b, c, ib, ic = chain
    assert ic.moved == [] and ic.stability_pct > 97
    assert ic.previous_live == len(b)                       # stability is measured against Rev B
    for k, slot in b.items():
        if k in c:
            assert c[k] == slot, k                           # everything alive in B and C kept its Rev B slot
    # a signal Rev B added is "unchanged" in C: it did not get re-placed by the Rev A baseline
    added_in_b = {m.tag.replace(" ", "") for m in ib.added}
    assert added_in_b & set(c) and all(c[k] == b[k] for k in added_in_b & set(c))


def test_rev_c_verifies_and_its_slots_carry_rev_numbers(chain):
    svc, st, *_ = chain
    assert st.current.report.verdict == "PASS" and st.current.report.violations == 0
    new = {s.revision for s in st.current.out.alloc.slots if s.change_type == "Added"}
    assert new == {"02"}
    kept_b = {s.revision for s in st.current.out.alloc.slots if s.status == "Field" and s.change_type in ("-", "", None) and s.revision != "00"}
    assert kept_b <= {"01"}                                  # a Rev B change is not re-stamped as Rev C


def test_switch_between_revisions_is_reproducible(chain):
    svc, st, *_ = chain
    h_c = st.current.cert.output_sha256
    switch(svc, "p", "A")
    assert st.project.revision == "A" and len(st.current.out.alloc.slots) == 1152
    switch(svc, "p", "B")
    assert st.current.impact.to_rev == "B"
    switch(svc, "p", "C")
    assert st.current.cert.output_sha256 == h_c


def test_chain_survives_a_restart(chain, tmp_path):
    svc, st, *_ = chain
    h = st.current.cert.output_sha256
    svc2 = Service(svc.store)
    st2 = svc2.get("p")
    assert st2.project.revision == "C" and list(svc2.revs(st2)) == ["B", "C"]
    assert st2.current.cert.output_sha256 == h


def test_cannot_skip_a_letter_and_replacing_drops_the_later_revisions(tmp_path):
    svc = fresh(tmp_path)
    with pytest.raises(ValueError, match="next revision is Rev B"):
        apply_revision(svc, "p", SAMPLE_BYTES, "x.xlsx", rev="D")
    run_gauntlet(svc, "p")
    run_gauntlet(svc, "p", rev="C")
    st = svc.get("p")
    assert list(svc.revs(st)) == ["B", "C"]
    run_gauntlet(svc, "p")                                   # rebuilding Rev B invalidates Rev C (pinned against the old B)
    assert list(svc.revs(st)) == ["B"] and st.project.revision == "B"
    with pytest.raises(ValueError):
        switch(svc, "p", "C")


def test_uploading_a_file_appends_the_next_revision(tmp_path):
    svc = fresh(tmp_path)

    def b_edit(ws):                      # delete one tag
        h = header(ws)
        ws.cell(5, h["Status"]).value = "DELETED"

    def c_edit(ws):                      # delete another one
        h = header(ws)
        ws.cell(9, h["Status"]).value = "DELETED"

    rb = edit(SAMPLE_BYTES, b_edit)
    apply_revision(svc, "p", rb, "revB.xlsx")
    rc = edit(rb, c_edit)
    ic = apply_revision(svc, "p", rc, "revC.xlsx")
    st = svc.get("p")
    assert st.project.revisions == ["A", "B", "C"]
    assert (ic.from_rev, ic.to_rev) == ("B", "C") and len(ic.deleted) == 1 and ic.moved == []
    assert ic.previous_live == 188                           # Rev B had 188 live signals, not Rev A's 189


def test_gauntlet_c_does_not_collide_with_rev_b_rows(chain):
    svc, st, a, b, c, ib, ic = chain
    tags_b = {m.tag.replace(" ", "") for m in ib.added}
    tags_c = {m.tag.replace(" ", "") for m in ic.added}
    assert not tags_b & tags_c


# ───────────────────────────── TQ Apply: pinned re-allocation ─────────────────────────────

def mdb_with_unmapped_cable() -> bytes:
    """The sample plus one signal on a NEW 24-pair cable whose IO type ('XX') is unknown. Mapping 'XX' makes the cable
    appear in the middle of the sort order, so a from-scratch re-allocation would shuffle the cables after it."""
    def add(ws):
        h = header(ws)
        src = next(r for r in range(2, ws.max_row + 1) if ws.cell(r, h["Multipair Cable No."]).value == "320-CE-0004")
        new = ws.max_row + 1
        for j in range(1, ws.max_column + 1):
            ws.cell(new, j).value = ws.cell(src, j).value
        ws.cell(new, h["Tag Number"]).value = "320-ZZ-99999"
        ws.cell(new, h["Multipair Cable No."]).value = "320-CE-0004A"
        ws.cell(new, h["Pair Number"]).value = "Pair 1"
        ws.cell(new, h["I/O Type Name"]).value = "XX"
    return edit(SAMPLE_BYTES, add)


def shuffler(tmp_path) -> Service:
    return fresh(tmp_path, mdb_with_unmapped_cable(), "unmapped.xlsx")


def tq(svc, category: str):
    return next(t for t in svc.get("p").current.tqs if t.category == category and t.status == "open")


def map_to_air(svc):
    t = tq(svc, "Unknown IO type")
    apply_tq(svc, "p", t.id, next(i for i, o in enumerate(t.fix_options) if o.label.endswith("→ AIR")))
    return t


def test_tq_apply_keeps_every_other_cable_in_place(tmp_path):
    svc = shuffler(tmp_path)
    st = svc.get("p")
    before = slots_of(st.current.out)
    t = map_to_air(svc)
    after = slots_of(st.current.out)
    assert len(after) == len(before) + 1 and "320-ZZ-99999" in after          # the fix allocated the new signal
    assert all(after[k] == v for k, v in before.items())                       # …and not one other signal moved
    assert st.current.report.violations == 0
    assert len(svc.tq_pins(st)) >= 1
    assert next(x for x in st.current.tqs if x.id == t.id).status == "applied"


def test_without_the_pins_the_same_fix_shuffles_cables(tmp_path):
    """Control: this fix really would move signals if it were re-allocated from scratch."""
    svc = shuffler(tmp_path)
    st = svc.get("p")
    before = slots_of(st.current.out)
    map_to_air(svc)
    svc.set_tq_pins(st, {})
    svc.recompute(st, record=False)
    scratch = slots_of(st.current.out)
    assert len([k for k, v in before.items() if scratch.get(k) != v]) > 20


def test_pins_survive_a_recompute_and_a_restart(tmp_path):
    svc = shuffler(tmp_path)
    st = svc.get("p")
    before = slots_of(st.current.out)
    map_to_air(svc)
    after, h = slots_of(st.current.out), st.current.cert.output_sha256
    svc.recompute(st, record=False)                                            # a plain re-allocate reproduces the layout
    assert slots_of(st.current.out) == after and st.current.cert.output_sha256 == h
    st2 = Service(svc.store).get("p")                                          # …and so does a restart
    assert slots_of(st2.current.out) == after and st2.current.cert.output_sha256 == h
    assert all(after[k] == v for k, v in before.items())


def test_pinned_apply_exports_a_verified_sheet(tmp_path):
    svc = shuffler(tmp_path)
    st = svc.get("p")
    map_to_air(svc)
    wb = openpyxl.load_workbook(io.BytesIO(st.current.xlsx))
    ws = wb["IO Assignment"]
    heads = {c.value: c.column for c in ws[1]}
    changes = {ws.cell(r, heads["Change Type"]).value for r in range(2, ws.max_row + 1)}
    assert changes <= {"-", None, ""}                                          # pins are silent: no 'Moved' noise in Rev A
    assert st.current.report.verdict == "PASS"


def test_a_hardware_only_fix_stores_no_pins(tmp_path):
    f = DRILLS / "drill_B_title_rows_blank_rows_junk.xlsx"
    svc = fresh(tmp_path, f.read_bytes(), f.name)
    t = tq(svc, "Unknown DO type")
    st = svc.get("p")
    before = slots_of(st.current.out)
    apply_tq(svc, "p", t.id, 0)                                                # treat UZV as high-current: charm/TB swap in place
    assert svc.tq_pins(st) == {}
    assert slots_of(st.current.out) == before
    uzv = {x.tag_key for x in st.current.out.signals if x.instrument_type == "UZV" and not x.quarantined}
    hw = {(s.charm, s.tb) for s in st.current.out.alloc.slots if s.tag_key in uzv and s.status == "Field"}
    from engine.hardware import listed_hw
    assert uzv and hw == {listed_hw(st.current.out.rules, "high_current_do")}


def test_tq_apply_in_a_revision_is_pinned_to_that_revision(tmp_path):
    svc = fresh(tmp_path)
    run_gauntlet(svc, "p")
    st = svc.get("p")
    before = slots_of(st.current.out)
    t = tq(svc, "Unknown DO type")
    apply_tq(svc, "p", t.id, 0)
    assert slots_of(st.current.out) == before and st.project.revision == "B"
    assert st.current.report.violations == 0 and st.current.impact.moved == []


# ───────────────────────────── TQ data patches ─────────────────────────────

def drill_b(tmp_path) -> Service:
    f = DRILLS / "drill_B_title_rows_blank_rows_junk.xlsx"
    return fresh(tmp_path, f.read_bytes(), f.name)


# ───────────────────────────── TQ data patches ─────────────────────────────

def test_missing_pair_tq_offers_and_applies_a_data_patch(tmp_path):
    svc = drill_b(tmp_path)
    st = svc.get("p")
    before = slots_of(st.current.out)
    t = tq(svc, "Missing pair")
    i = next(i for i, o in enumerate(t.fix_options) if o.data_patch)
    patch = t.fix_options[i].data_patch
    assert patch["fields"]["pair"].isdigit() and patch["tag"] == t.tags[0]
    apply_tq(svc, "p", t.id, i)
    assert t.tags[0].replace(" ", "") in slots_of(st.current.out)           # the quarantined row is now allocated
    after = slots_of(st.current.out)
    assert all(after[k] == v for k, v in before.items())                    # and nothing else moved
    assert st.current.report.violations == 0
    from app.signalfix import list_patches
    assert any(p.tag_key == t.tags[0].replace(" ", "") for p in list_patches(svc, "p"))
    log = svc.store.log("decisions", "p")
    assert any(d.get("tq_id") == t.id and d.get("who") == "TQ Apply" for d in log)
    assert next(x for x in st.current.tqs if x.id == t.id).status == "applied"


def test_duplicate_tag_tq_deletes_the_duplicate_row(tmp_path):
    svc = drill_b(tmp_path)
    st = svc.get("p")
    t = tq(svc, "Duplicate tag")
    i = next(i for i, o in enumerate(t.fix_options) if o.data_patch)
    apply_tq(svc, "p", t.id, i)
    assert not [x for x in st.current.tqs if x.category == "Duplicate tag" and x.status == "open"]
    assert st.current.report.violations == 0


def test_data_patch_tq_apply_is_undoable(tmp_path):
    from app.sandbox import undo
    svc = drill_b(tmp_path)
    st = svc.get("p")
    n0 = len(slots_of(st.current.out))
    t = tq(svc, "Missing pair")
    apply_tq(svc, "p", t.id, next(i for i, o in enumerate(t.fix_options) if o.data_patch))
    assert len(slots_of(st.current.out)) == n0 + 1
    undo(svc, "p")
    assert len(slots_of(st.current.out)) == n0


def test_a_fix_that_needs_a_value_is_refused_with_a_reason(tmp_path):
    from app.schemas import TQ
    from app.sandbox import _data_patch_target
    tq_ = TQ(id="TQ-X-001", category="c", severity="low", message="m", tags=["T-1"], rows=[3])
    with pytest.raises(ValueError, match="needs a value"):
        _data_patch_target(tq_, {"field": "tag"})
    assert _data_patch_target(tq_, {"field": "pair", "value": 4}) == ("T-1", 3, {"pair": "4"})
    assert _data_patch_target(tq_, {"tag": "T-2", "row": 9, "fields": {"pair": "5"}}) == ("T-2", 9, {"pair": "5"})
