"""Make a tampered copy of an exported IO Allocation workbook — the Proof Center demo.

    uv run --project backend python scripts/make_tampered.py <exported.xlsx> <tampered.xlsx> [--nth 40]

Changes ONE cell: the CHARM Type of the n-th Field row becomes a charm that does not match the hardware map.
Dropped into Proof Center → "Verify any Excel", it fails the hardware-map check at that exact Excel row.
"""
import argparse
import sys

import openpyxl

WRONG = "SE4303T03"  # RTD input charm: wrong for any non-RTD signal


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("src")
    ap.add_argument("dst")
    ap.add_argument("--nth", type=int, default=40, help="which Field row to tamper (1-based)")
    args = ap.parse_args()

    wb = openpyxl.load_workbook(args.src)
    ws = wb["IO Assignment"] if "IO Assignment" in wb.sheetnames else wb.active
    col = {str(c.value).strip(): c.column for c in ws[1] if c.value}
    charm, status, tag = col["CHARM Type"], col["Tag Status"], col["Tag Name"]
    n = 0
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, status).value == "Field":
            n += 1
            if n == args.nth:
                old = ws.cell(r, charm).value
                ws.cell(r, charm).value = WRONG if old != WRONG else "SE4304T01"
                wb.save(args.dst)
                print(f"tampered row {r}: {ws.cell(r, tag).value!r} CHARM Type {old} -> {ws.cell(r, charm).value}")
                return 0
    print(f"fewer than {args.nth} Field rows", file=sys.stderr)
    return 1


if __name__ == "__main__":
    raise SystemExit(main())
