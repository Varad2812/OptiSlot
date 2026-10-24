"""Extract the hardware catalogue (CHARM / TB part numbers and descriptions) and the
output template spec from the reference workbook and the blank template.

Writes backend/catalog.json. Descriptions are copied exactly as spelled in the reference.
Usage: uv run --project backend python scripts/extract_catalog.py
"""
from __future__ import annotations

import json
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
REF = ROOT / "data" / "Sample_IO_Allocation_VIT_v2_1.xlsx"
TEMPLATE = ROOT / "data" / "Blank_IO_Allocation_Template.xlsx"
OUT = ROOT / "backend" / "catalog.json"


def _rows(path: Path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    it = ws.iter_rows(values_only=True)
    header = [str(h) if h is not None else "" for h in next(it)]
    for row in it:
        yield dict(zip(header, ["" if v is None else str(v).strip() for v in row]))


def main() -> None:
    charms: dict[str, Counter] = defaultdict(Counter)
    tbs: dict[str, Counter] = defaultdict(Counter)
    spare: Counter = Counter()
    for r in _rows(REF):
        if r.get("Tag Status") == "Spare Slot":
            spare[(r.get("CHARM Type", "-"), r.get("CHARM TB Part No", "-"))] += 1
        if r.get("CHARM Type") and r["CHARM Type"] != "-":
            charms[r["CHARM Type"]][r.get("Charm Description", "")] += 1
        if r.get("CHARM TB Part No") and r["CHARM TB Part No"] != "-":
            tbs[r["CHARM TB Part No"]][r.get("CHARM TB Description", "")] += 1

    wb = openpyxl.load_workbook(TEMPLATE)
    ws = wb.worksheets[0]
    columns = [str(c.value) for c in ws[1]]

    catalog = {
        "source": REF.name,
        "charms": {k: v.most_common(1)[0][0] for k, v in sorted(charms.items())},
        "terminal_blocks": {k: v.most_common(1)[0][0] for k, v in sorted(tbs.items())},
        "spare_slot": dict(zip(("charm", "tb"), spare.most_common(1)[0][0])),   # hardware of a slot with no cable
        "template": {
            "file": TEMPLATE.name,
            "sheet": ws.title,
            "columns": columns,
            "preformatted_rows": ws.max_row,
            "header_fill": "1F3864",
        },
    }
    OUT.write_text(json.dumps(catalog, indent=2, ensure_ascii=False) + "\n")
    print(f"wrote {OUT}: {len(catalog['charms'])} CHARMs, {len(catalog['terminal_blocks'])} TBs, {len(columns)} columns")
    for k, v in catalog["charms"].items():
        print(f"  {k:10} {v}")
    for k, v in catalog["terminal_blocks"].items():
        print(f"  {k:10} {v}")


if __name__ == "__main__":
    main()
